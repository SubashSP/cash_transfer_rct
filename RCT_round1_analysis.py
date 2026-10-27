"""
Cash transfer RCT — Survey Round 1: data merging and analysis (Python)
=====================================================================

Reproduces, in one script, everything done in this session:

  Step 1  merge_survey()   Merge the 20 repeat sheets of Survey_Round_1.xlsx into the main sheet
                           (long -> wide, older and newer form versions stacked), correct two farmer
                           IDs (165 -> 166 for _index 234; 201 -> 200 for _index 16), add `round` = 1,
                           `payment` and `in_payment_list` from RCT_payment_round1.xls, compute
                           Section 3 activity costs, costs per acre and outlier flags, and write
                           Survey_Round_1_merged_wide.xlsx (sheets: data, codebook, data_checks,
                           cost_summary, cost_outliers, cost_review).
  Step 2  balance_table()  Baseline balance of socio-economic characteristics by payment
                           -> Balance_table_payment_round1.xlsx
  Step 3  cost_comparison() Per-acre cost components, treatment vs control
                           -> Cost_comparison_payment_round1.xlsx
  Step 4  timeline_items() Item-level costs with activity / payment dates, put on the 2025 calendar,
                           weekly aggregates (also the data behind the interactive timeline page)
  Step 5  bubble_plot()    Weekly cost-per-acre bubble plot with the cumulative cost per acre of the
                           average farmer -> Cost_per_acre_bubble_overlay.jpg / _panels.jpg

Usage
-----
  Put the three input files in one folder (DATA_DIR below) and run:
      python RCT_round1_analysis.py            # uses DATA_DIR / OUT_DIR below
      python RCT_round1_analysis.py <data_dir> <out_dir>

Requirements: pandas, numpy, scipy, openpyxl, matplotlib, and xlrd (to read the .xls payment file:
`pip install xlrd`). If xlrd is not available, save RCT_payment_round1.xls as .xlsx in the same folder.
"""
import json
import os
import re
import sys
import textwrap
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.lines import Line2D
from matplotlib.ticker import FixedLocator, FuncFormatter, NullLocator
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# =====================================================================================
# CONFIGURATION — edit these two paths
# =====================================================================================
DATA_DIR = Path(".")            # folder with Survey_Round_1.xlsx, Survey_Round_1_questionnarie.xlsx, RCT_payment_round1.xls
OUT_DIR = Path("outputs")       # all outputs are written here


def read_payment(path):
    """Read RCT_payment_round1.xls (needs xlrd); fall back to an .xlsx copy with the same name."""
    path = Path(path)
    try:
        return pd.read_excel(path)
    except (ImportError, ValueError, FileNotFoundError) as err:
        alt = path.with_suffix(".xlsx")
        if alt.exists():
            return pd.read_excel(alt)
        raise RuntimeError(f"Could not read {path} ({err}). Install xlrd (pip install xlrd) "
                           f"or save the file as {alt.name}.") from err


# =====================================================================================
# SECTION 3 COST RULES (activity costs, per-acre costs, outlier screen)
# =====================================================================================
CODES = [999, 9999]                 # "don't know" codes used by enumerators
LUMP = re.compile(r"contract|contact|lump|total", re.I)
SMALL_UNIT = re.compile(r"\bml\b|\bg\b|gm|gram", re.I)
# chemical fertilisers and lime are sold per kg / bag at well above Rs 3 per unit
PER_KG_ITEM = re.compile(r"urea|potash|\bmop\b|fact|dap|sulph|phos|npk|complex|ammon|micro|magnes|zinc|borax|lime|line|dolomite", re.I)
FLOOR_ANY = 1.0       # amount ÷ quantity below Rs 1 per unit: cannot be a total for any item → amount is a unit price
FLOOR_FERT = 3.0      # for fertiliser / lime items the floor is Rs 3 per unit

# (key, label, section, qty, unit, price)  — price is per unit in the questionnaire
FIXED = [
    ("a_seed", "Seed", "A", "a_seed_qty", "a_seed_unit", "a_seed_price"),
    ("a_land", "Land preparation", "A", "a_land_qty", "a_land_unit", "a_land_price"),
    ("b_labour", "Nursery labour", "B", "b_labour_qty", "b_labour_unit", "b_labour_amt"),
    ("b_fert", "Nursery fertiliser", "B", "b_fert_qty", "b_fert_unit", "b_fert_amt"),
    ("c_mainlab", "Main-plot preparation labour", "C", "c_labour_qty", "c_labour_unit", "c_labour_amt"),
    ("c_mach", "Main-plot preparation machinery", "C", "c_mach_qty", "c_mach_unit", "c_mach_amt"),
    ("c_translab", "Transplanting labour", "C", "c_labour_qty_001", "c_labour_unit_001", "c_labour_amt_001"),
    ("c_fert1lab", "Labour, 1st fertiliser dose", "C", "c_labour_fert1_qty", "c_labour_fert1_unit", "c_labour_fert1_amt"),
    ("c_gapfill", "Gap-filling labour", "C", "c_labour_gapfilling_qty", "c_labour_gapfilling_unit", "c_labour_gapfilling_amt"),
    ("d_weed", "Herbicide / weedicide", "D", "d_weed_qty", "d_weed_unit", "d_weed_price"),
    ("d_weedlab", "Labour, weedicide application", "D", "d_weedlab_qty", "d_weedlab_unit", "d_weedlab_price"),
    ("d_fert2lab", "Labour, 2nd fertiliser dose", "D", "d_labour_fert2_qty", "d_labour_fert2_unit", "d_labour_fert2_amt"),
    ("e_pest", "Pesticide", "E", "e_pest_qty", "e_pest_unit", "e_pest_price"),
    ("e_pestlab", "Labour, pesticide application", "E", "e_labour_pest_qty", "e_labour_pest_unit", "e_labour_pest_amt"),
    ("e_fert3lab", "Labour, 3rd fertiliser dose", "E", "e_labour_fert3_qty", "e_labour_fert3_unit", "e_labour_fert3_amt"),
    ("f_mach", "Harvesting machinery", "F", "f_mach_qty", "f_mach_unit", "f_mach_amt"),
    ("f_tractor", "Tractor, harvest transport", "F", "f_tractor_qty", "f_tractor_unit", "f_tractor_amt"),
    ("f_labour", "Labour, drying and packing", "F", "f_labour_qty", "f_labour_unit", "f_labour_amt"),
]
# repeat blocks with a per-unit price: (key, label, section, prefix for qty/unit/price, max items)
FERT = [
    ("c_fert1", "Fertiliser, 1st dose", "C", "c_fert1", 2),
    ("d_fert2", "Fertiliser, 2nd dose", "D", "d_fert2", 6),
    ("e_fert3", "Fertiliser, 3rd dose / pesticide", "E", "e_fert3", 5),
]
# repeat blocks asking "amount paid (Rs.)": (key, label, section, prefix, max items)
OTHER = [
    ("a_other", "Other items, pre-nursery", "A", "a_other", 4),
    ("b_other", "Other items, nursery", "B", "b_other", 4),
    ("c_other", "Other items, after transplanting", "C", "c_other", 5),
    ("d_other", "Other items, after 2nd dose", "D", "d_other", 3),
    ("e_other", "Other items, pesticide / 4th dose", "E", "e_other", 6),
    ("f_other", "Other items, after harvest", "F", "f_other", 4),
]
ORDER = ["a_seed", "a_land", "a_other", "b_labour", "b_fert", "b_other", "c_mainlab", "c_mach", "c_translab", "c_fert1",
         "c_fert1lab", "c_gapfill", "c_other", "d_weed", "d_weedlab", "d_fert2", "d_fert2lab", "d_other",
         "e_pest", "e_pestlab", "e_fert3", "e_fert3lab", "e_other", "f_mach", "f_tractor", "f_labour", "f_other"]
SECTIONS = {"A": "Pre-nursery preparation", "B": "Nursery preparation", "C": "Transplanting",
            "D": "Weedicide / fertiliser application", "E": "Pesticide / fertiliser application", "F": "Harvesting"}


def _num(s):
    return pd.to_numeric(s, errors="coerce")


def unit_price_cost(q, unit, p):
    """cost = quantity × unit price, row-wise. Returns (cost, status)."""
    q, p = _num(q), _num(p)
    unit = unit.astype(str).where(unit.notna(), "")
    cost = pd.Series(np.nan, index=q.index)
    status = pd.Series("", index=q.index, dtype=object)
    code = q.isin(CODES) | p.isin(CODES)
    both_blank = q.isna() & p.isna()
    zero = (q == 0) | (p == 0)
    lump = q.isna() & (p > 0) & unit.str.contains(LUMP)
    ok = q.notna() & p.notna()
    cost[ok] = q[ok] * p[ok]
    cost[lump] = p[lump]; status[lump] = "quantity blank, unit is contract → price taken as lump sum"
    cost[zero & ~code] = 0.0
    cost[both_blank] = 0.0
    cost[code] = np.nan; status[code] = "999 / 9999 code in quantity or price → missing"
    half = cost.isna() & ~code
    status[half & q.isna()] = "price recorded but quantity blank → missing"
    status[half & p.isna()] = "quantity recorded but price blank → missing"
    return cost, status


def other_cost(q, unit, a, item):
    """'Amount paid (Rs.)' items: amount is the total paid, unless it is evidently a unit price."""
    q, a = _num(q), _num(a)
    unit = unit.astype(str).where(unit.notna(), "")
    item = item.astype(str).where(item.notna(), "")
    qq = q.where(~q.isin(CODES))
    cost = a.copy().astype(float)
    status = pd.Series("", index=q.index, dtype=object)
    cost[a.isna() & qq.isna()] = 0.0
    cost[a == 0] = 0.0
    ratio = a / qq
    base = (qq > 1) & (a > 0) & ~unit.str.contains(SMALL_UNIT) & ~a.isin(CODES)
    r1 = base & (ratio < FLOOR_ANY)
    r2 = base & ~r1 & (ratio < FLOOR_FERT) & item.str.contains(PER_KG_ITEM)
    amb = base & ~r1 & ~r2 & (ratio < FLOOR_FERT)
    cost[r1 | r2] = (qq * a)[r1 | r2]
    status[r1] = f"amount ÷ quantity < Rs {FLOOR_ANY:g} per unit → amount is a unit price: cost = quantity × amount"
    status[r2] = f"fertiliser/lime item, amount ÷ quantity < Rs {FLOOR_FERT:g} per unit → amount is a unit price: cost = quantity × amount"
    status[amb] = f"ambiguous (amount ÷ quantity between Rs {FLOOR_ANY:g} and {FLOOR_FERT:g}) → kept as total paid; please check"
    code = a.isin(CODES)
    cost[code] = np.nan; status[code] = "999 / 9999 code in amount → missing"
    miss = a.isna() & (qq > 0)
    cost[miss] = np.nan; status[miss] = "quantity recorded but amount blank → missing"
    return cost, status


def fert_items(d, pre, n):
    """(qty, unit, price, name) columns of a fertiliser block: repeat items 1..n, plus the single main-sheet item
    that the older Round 1 form used for the 1st dose (c_fert1_qty, c_fert1_price, ...), where present."""
    items = [(f"{pre}_qty_{k}", f"{pre}_unit_{k}", f"{pre}_price_{k}", f"{pre}_name_{k}") for k in range(1, n + 1)]
    if all(c in d.columns for c in (f"{pre}_qty", f"{pre}_unit", f"{pre}_price", f"{pre}_name")):
        items.append((f"{pre}_qty", f"{pre}_unit", f"{pre}_price", f"{pre}_name"))
    return items


def compute(d):
    """Return (costs DataFrame aligned to d, review rows, labels dict)."""
    out = pd.DataFrame(index=d.index)
    incomplete = pd.DataFrame(index=d.index)
    review = []
    labels = {}
    acres = _num(d["rice_acres_2025"])

    def log(key, label, status, extra):
        for i in status.index[status != ""]:
            review.append({"farmers_id": d.at[i, "farmers_id"], "_index": d.at[i, "_index"], "activity": key,
                           "activity_label": label, **{k: v.at[i] if hasattr(v, "at") else v for k, v in extra.items()},
                           "decision": status.at[i]})

    for key, label, sec, q, u, p in FIXED:
        c, st = unit_price_cost(d[q], d[u], d[p])
        out[f"cost_{key}"] = c
        incomplete[key] = c.isna()
        labels[key] = (label, sec, f"{q} × {p}")
        log(key, label, st, {"item": "", "quantity": d[q], "unit": d[u], "price_or_amount": d[p], "cost_used": c})

    for key, label, sec, pre, n in FERT:
        tot = pd.Series(0.0, index=d.index)
        anymiss = pd.Series(False, index=d.index)
        for q, u, p, nm in fert_items(d, pre, n):
            c, st = unit_price_cost(d[q], d[u], d[p])
            anymiss |= c.isna()
            tot = tot + c.fillna(0)
            log(key, label, st, {"item": d[nm], "quantity": d[q], "unit": d[u], "price_or_amount": d[p], "cost_used": c})
        out[f"cost_{key}"] = tot.where(~anymiss | (tot > 0))      # all items unusable → missing
        incomplete[key] = anymiss
        single = f" + {pre}_qty × {pre}_price (older form: one item in the main sheet)" if f"{pre}_qty" in d.columns else ""
        labels[key] = (label, sec, f"Σ items ({pre}_qty_k × {pre}_price_k){single}")

    for key, label, sec, pre, n in OTHER:
        tot = pd.Series(0.0, index=d.index)
        anymiss = pd.Series(False, index=d.index)
        for k in range(1, n + 1):
            q, u, a = f"{pre}_qty_{k}", f"{pre}_unit_{k}", f"{pre}_amt_{k}"
            c, st = other_cost(d[q], d[u], d[a], d[f"{pre}_item_{k}"])
            anymiss |= c.isna()
            tot = tot + c.fillna(0)
            log(key, label, st, {"item": d[f"{pre}_item_{k}"], "quantity": d[q], "unit": d[u], "price_or_amount": d[a], "cost_used": c})
        out[f"cost_{key}"] = tot.where(~anymiss | (tot > 0))
        incomplete[key] = anymiss
        labels[key] = (label, sec, f"Σ items: {pre}_amt_k (total paid), or {pre}_qty_k × {pre}_amt_k where the amount is evidently a unit price")

    out = out[[f"cost_{k}" for k in ORDER]]
    n_missing = incomplete.sum(axis=1)
    for s in SECTIONS:
        cols = [f"cost_{k}" for k in ORDER if labels[k][1] == s]
        out[f"cost_sec{s}"] = out[cols].sum(axis=1, min_count=1)
    out["cost_total"] = out[[f"cost_sec{s}" for s in SECTIONS]].sum(axis=1, min_count=1)
    out["cost_n_missing"] = n_missing.astype(int)
    valid_acres = acres > 0
    for c in [f"cost_{k}" for k in ORDER] + [f"cost_sec{s}" for s in SECTIONS] + ["cost_total"]:
        out["cpa_" + c[5:]] = (out[c] / acres).where(valid_acres)
    return out, pd.DataFrame(review), labels


def modified_z_log(s):
    """Iglewicz–Hoaglin modified z-score on log(value), positive values only."""
    x = np.log(s[s > 0])
    med = x.median()
    mad = (x - med).abs().median()
    z = pd.Series(np.nan, index=s.index)
    if mad > 0:
        z[x.index] = 0.6745 * (x - med) / mad
    return z


def _fmt(x):
    return f"{x:,.0f}" if abs(x) >= 10 else f"{x:,.2f}"


def _items(d, k):
    """Item-level (qty, unit, price/amount, name) frames for an activity, stacked over repeat positions."""
    for key, label, sec, q, u, p in FIXED:
        if key == k:
            return [pd.DataFrame({"qty": _num(d[q]), "unit": d[u], "price": _num(d[p]), "name": ""})]
    for key, label, sec, pre, n in FERT:
        if key == k:
            return [pd.DataFrame({"qty": _num(d[q]), "unit": d[u], "price": _num(p_), "name": d[nm]})
                    for q, u, p_, nm in [(q, u, d[p], nm) for q, u, p, nm in fert_items(d, pre, n)]]
    for key, label, sec, pre, n in OTHER:
        if key == k:
            return [pd.DataFrame({"qty": _num(d[f"{pre}_qty_{j}"]), "unit": d[f"{pre}_unit_{j}"],
                                  "price": _num(d[f"{pre}_amt_{j}"]), "name": d[f"{pre}_item_{j}"]}) for j in range(1, n + 1)]
    return None


def outliers(d, cst, labels, zcut=3.5):
    acres = _num(d["rice_acres_2025"])
    keys = ORDER + [f"sec{s}" for s in SECTIONS] + ["total"]
    rows, summ = [], []
    flags = pd.DataFrame(index=d.index)
    zs = {}
    for k in keys:
        zs[k] = modified_z_log(cst[f"cpa_{k}"])
        flags[f"out_{k}"] = (zs[k].abs() > zcut).astype(int)

    def acre_note(i):
        a = acres.at[i]
        if a <= 0.5:
            return f"very small rice area ({a:g} acre) inflates the per-acre figure — check rice_acres_2025"
        if a > 20:
            return f"very large rice area ({a:g} acres) for the costs reported — check rice_acres_2025"
        return None

    def cause(i, k, z):
        hi = z > 0
        if k.startswith("sec") or k == "total":
            sec_keys = [a for a in ORDER if k == "total" or labels[a][1] == k[3:]]
            drv = [a for a in sec_keys if flags.at[i, f"out_{a}"] and np.sign(zs[a].get(i, 0)) == np.sign(z)]
            if drv:
                return "driven by: " + "; ".join(f"{labels[a][0]} (Rs {_fmt(cst.at[i, f'cpa_{a}'])}/acre)" for a in drv)
            return acre_note(i) or "no single activity flagged — several moderately unusual items add up"
        parts = _items(d, k)
        allit = pd.concat(parts)
        mq = allit.loc[allit["qty"] > 0, "qty"].median()
        mp = allit.loc[allit["price"] > 0, "price"].median()
        cand = pd.DataFrame([{"qty": f.at[i, "qty"], "unit": f.at[i, "unit"], "price": f.at[i, "price"], "name": f.at[i, "name"]}
                             for f in parts])
        cand["c"] = cand["qty"].fillna(1) * cand["price"].fillna(0)
        r = cand.loc[cand["c"].idxmax()] if hi else cand.loc[cand["price"].fillna(np.inf).idxmin()]
        nm = f"{str(r['name']).strip()}: " if str(r["name"]).strip() not in ("", "nan") else ""
        qr = r["qty"] / mq if pd.notna(r["qty"]) and mq else np.nan
        pr = r["price"] / mp if pd.notna(r["price"]) and mp else np.nan
        unit = "" if pd.isna(r["unit"]) else str(r["unit"]).strip()
        desc = f"{nm}qty {_fmt(r['qty']) if pd.notna(r['qty']) else '—'} {unit} × Rs {_fmt(r['price']) if pd.notna(r['price']) else '—'}".replace("  ", " ")
        what = "amount" if k in [o[0] for o in OTHER] else "unit price"
        if hi:
            if re.search(r"\bkg", unit, re.I) and pd.notna(r["price"]) and r["price"] > 300 and what == "unit price":
                return (f"{desc}. Rs {_fmt(r['price'])} per kg is implausible for this input — looks like the bag price or the total "
                        f"was entered as the unit price")
            if qr >= 20:
                return f"{desc}. Quantity is {qr:,.0f}× the median ({_fmt(mq)}) — looks like the total amount was entered as quantity"
            if pr >= 10:
                return f"{desc}. {what.capitalize()} is {pr:,.0f}× the median ({_fmt(mp)}) — looks like the total was entered as {what}"
            a = acre_note(i)
            if a:
                return f"{desc}. {a}"
            if (pd.isna(qr) or qr < 2) and (pd.isna(pr) or pr < 2):
                return f"{desc}. Quantity and {what} are typical; per-acre value is high relative to the {acres.at[i]:g}-acre rice area — check rice_acres_2025"
            return f"{desc}. High quantity ({qr:,.1f}× median) and/or {what} ({pr:,.1f}× median) — check"
        a = acre_note(i)
        if a and acres.at[i] > 20:
            return f"{desc}. {a}"
        if (pd.isna(qr) or qr >= 0.5) and (pd.isna(pr) or pr >= 0.5):
            return (f"{desc}. Quantity and {what} are typical; per-acre value is low for a {acres.at[i]:g}-acre rice area — "
                    f"input may cover only part of the area, or check rice_acres_2025")
        return f"{desc}. Unusually low (qty {qr:,.2f}×, {what} {pr:,.2f}× median) — check units, a partial payment or own/family labour"

    for k in keys:
        cpa = cst[f"cpa_{k}"]
        pos = cpa[cpa > 0]
        z = zs[k]
        q1, q3 = pos.quantile([0.25, 0.75]) if len(pos) else (np.nan, np.nan)
        hi, lo = z > zcut, z < -zcut
        if k in labels:
            lab, sec = labels[k][0], labels[k][1]
        elif k.startswith("sec"):
            lab, sec = f"Section {k[3:]} total: {SECTIONS[k[3:]]}", k[3:]
        else:
            lab, sec = "TOTAL cost of cultivation", "All"
        lx = np.log(pos); mad = (lx - lx.median()).abs().median()
        summ.append({"section": sec, "activity": lab, "variable": f"cpa_{k}",
                     "N with cost > 0": len(pos), "N cost = 0": int((cpa == 0).sum()), "N missing": int(cpa.isna().sum()),
                     "Mean": pos.mean(), "Median": pos.median(), "P5": pos.quantile(0.05), "P95": pos.quantile(0.95),
                     "Min": pos.min(), "Max": pos.max(),
                     "Lower bound": float(np.exp(lx.median() - zcut * mad / 0.6745)),
                     "Upper bound": float(np.exp(lx.median() + zcut * mad / 0.6745)),
                     "N high outliers": int(hi.sum()), "N low outliers": int(lo.sum()),
                     "N > Q3 + 3×IQR": int((pos > q3 + 3 * (q3 - q1)).sum())})
        for i in z.index[hi | lo]:
            rows.append({"farmers_id": d.at[i, "farmers_id"], "_index": d.at[i, "_index"], "section": sec, "activity": lab,
                         "variable": f"cpa_{k}", "direction": "high" if z.at[i] > 0 else "low",
                         "rice_acres_2025": acres.at[i], "cost (Rs.)": cst.at[i, f"cost_{k}"], "cost per acre (Rs.)": cpa.at[i],
                         "median cost per acre": pos.median(), "times the median": cpa.at[i] / pos.median(),
                         "modified z (log)": z.at[i], "likely cause": cause(i, k, z.at[i])})
    # households with no usable rice area have no per-acre value: screen their raw activity cost instead
    no_area = ~(acres > 0)
    for k in ORDER:
        zc = modified_z_log(cst[f"cost_{k}"])
        pos = cst.loc[cst[f"cost_{k}"] > 0, f"cost_{k}"]
        for i in zc.index[(zc > zcut) & no_area]:
            flags.at[i, f"out_{k}"] = 1
            lab, sec = labels[k][0], labels[k][1]
            rows.append({"farmers_id": d.at[i, "farmers_id"], "_index": d.at[i, "_index"], "section": sec, "activity": lab,
                         "variable": f"cpa_{k}", "direction": "high",
                         "rice_acres_2025": acres.at[i], "cost (Rs.)": cst.at[i, f"cost_{k}"], "cost per acre (Rs.)": np.nan,
                         "median cost per acre": np.nan, "times the median": cst.at[i, f"cost_{k}"] / pos.median(),
                         "modified z (log)": zc.at[i],
                         "likely cause": "No rice area recorded, so screened on total cost (times the median = vs median total cost). "
                                         + cause(i, k, zc.at[i]).split(". very small rice area")[0]})
    for k in ORDER:                                    # keep the summary counts in step with the added rows
        extra = sum(1 for r in rows if r["variable"] == f"cpa_{k}" and pd.isna(r["cost per acre (Rs.)"]))
        for srow in summ:
            if srow["variable"] == f"cpa_{k}":
                srow["N high outliers"] += extra
    return pd.DataFrame(summ), pd.DataFrame(rows), flags

C = sys.modules[__name__]     # lets the step functions refer to the cost rules as C.<name>


# =====================================================================================
# STEP 1 — merge sheets, payment, IDs, costs  ->  Survey_Round_1_merged_wide.xlsx
# =====================================================================================
def merge_survey(IN, OUT):
    IN = Path(IN)
    SURVEY = IN / "Survey_Round_1.xlsx"
    QUEST = IN / "Survey_Round_1_questionnarie.xlsx"
    PAY = IN / "RCT_payment_round1.xls"
    MAIN_SHEET = "Farmer Survey - Round 1"

    # block name -> (count variable that triggers the repeat, [sheets: older version first, newer second], position var, section label)
    BLOCKS = [
        ("offfarm", "offfarm_count", ["offfarm_rep"], "mem_pos", "Sec 1: off-farm employed members"),
        ("crop", "other_crops_count", ["crop_rep"], "crop_pos", "Sec 2: other crops cultivated"),
        ("a_other", "a_other_count", ["begin_repeat_OOtl5mSOZ", "begin_repeat_0W8hGwBOT"], "a_other_pos", "Sec 3A: other items, pre-nursery"),
        ("b_other", "b_other_count", ["begin_repeat_FXOFif2yY", "begin_repeat_oFXmOdKGR"], "b_other_pos", "Sec 3B: other items, nursery"),
        ("d_fert1", "d_fert1_count", ["begin_repeat_q7UBsOtMr"], "d_fert1_pos", "Sec 3C: first-dose fertilisers (newer form version only)"),
        ("c_other", "c_other_count", ["begin_repeat_FV3YMN63y", "begin_repeat_kb3NZEaME"], "c_other_pos", "Sec 3C: other items, transplanting"),
        ("d_fert2", "d_fert2_count", ["begin_repeat_BgK7Aa7Uw", "begin_repeat_pHGxYqKsw"], "d_fert2_pos", "Sec 3D: second-dose fertilisers"),
        ("d_other", "d_other_count", ["begin_repeat_bDKIjSEX3", "begin_repeat_9aGJkPAfe"], "d_other_pos", "Sec 3D: other items, after 2nd dose"),
        ("e_fert3", "e_fert3_count", ["begin_repeat_ntYu79uDD", "begin_repeat_y8vUlnzPY"], "e_fert3_pos", "Sec 3E: pesticides / third-dose fertilisers"),
        ("e_other", "e_other_count", ["begin_repeat_z8dV2uORP", "begin_repeat_PS5svQ16g"], "e_other_pos", "Sec 3E: other items"),
        ("f_other", "f_other_count", ["begin_repeat_lM3hqngUR", "begin_repeat_tXgwButw2"], "f_other_pos", "Sec 3F: other items, harvesting"),
        ("welfare", "welfare_count", ["welfare_rep"], "scheme_pos", "Sec 6: welfare schemes"),
    ]

    xl = pd.ExcelFile(SURVEY)
    main = pd.read_excel(xl, sheet_name=MAIN_SHEET)
    n_main = len(main)
    assert main["_index"].is_unique

    # farmer ID corrections requested by the user: {_index: (old id, new id)}
    ID_FIX = {234: ("165", "166"),   # second submission of 165
              16: ("201", "200")}    # first submission of 201
    fix_notes = []
    for idx, (old, new) in ID_FIX.items():
        row = main.index[main["_index"] == idx][0]
        assert str(main.at[row, "farmers_id"]).strip() == old, f"_index {idx} is not {old}"
        assert new not in main["farmers_id"].astype(str).str.strip().values, f"{new} already in survey"
        main.at[row, "farmers_id"] = new
        fix_notes.append(f"_index {idx} (interview {main.at[row, 'interview_date']:%Y-%m-%d %H:%M}): {old} → {new}")

    # every repeat sheet must be used exactly once
    used = [s for b in BLOCKS for s in b[2]]
    assert sorted(used) == sorted(s for s in xl.sheet_names if s != MAIN_SHEET), "sheet list mismatch"

    # ---- questionnaire labels for the codebook ----
    q = pd.read_excel(QUEST, sheet_name="survey")
    qlab = {}
    for _, r in q.iterrows():
        if pd.notna(r["name"]):
            lab = r["label::English (en)"] if pd.notna(r["label::English (en)"]) else r.get("label")
            qlab[r["name"]] = re.sub(r"\s+", " ", str(lab)).strip() if pd.notna(lab) else ""
    qtype = dict(zip(q["name"], q["type"]))
    ch = pd.read_excel(QUEST, sheet_name="choices")
    choice_lab = {(r["list_name"], str(r["name"])): r["label::English (en)"] for _, r in ch.iterrows()}

    def label_for(var):
        if var in qlab:
            return qlab[var]
        if "/" in var:                       # select_multiple dummies, e.g. has_phone/1
            base, opt = var.split("/", 1)
            lab = qlab.get(base, "")
            t = str(qtype.get(base, ""))
            lst = t.split()[-1] if t.startswith("select_multiple") else None
            opt_lab = choice_lab.get((lst, opt)) if lst else None
            if lab:
                return f"{lab} — option {opt}" + (f": {opt_lab}" if opt_lab else "")
        return "(not in the attached questionnaire version)"

    # ---- reshape each block long -> wide ----
    wide_blocks = {}        # count_var -> wide df indexed by _parent_index
    codebook_extra = {}     # new column -> (source sheets, block, item, base var)
    checks = []
    for name, count_var, sheets, pos_var, section in BLOCKS:
        parts = []
        for s in sheets:
            d = pd.read_excel(xl, sheet_name=s)
            d["__src"] = s
            parts.append(d)
        long = pd.concat(parts, ignore_index=True)
        # sanity: parents exist, positions match the repeat order, no parent spread across both versions
        assert long["_parent_index"].isin(main["_index"]).all(), f"{name}: orphan rows"
        assert (long[pos_var] == long.groupby("_parent_index").cumcount() + 1).all(), f"{name}: position mismatch"
        assert long.groupby("_parent_index")["__src"].nunique().max() == 1, f"{name}: parent in both versions"
        datavars = [c for c in long.columns if not c.startswith("_") and c != pos_var]
        w = long.pivot(index="_parent_index", columns=pos_var, values=datavars)
        max_item = int(long[pos_var].max())
        cols = []
        for k in range(1, max_item + 1):          # item-major: all fields of item 1, then item 2, ...
            for v in datavars:
                cols.append((v, k))
        w = w.reindex(columns=pd.MultiIndex.from_tuples(cols))
        w.columns = [f"{v}_{k}" for v, k in cols]
        wide_blocks[count_var] = w
        for v, k in cols:
            codebook_extra[f"{v}_{k}"] = (" + ".join(sheets), section, k, v)

        # classify repeat rows: blank (all fields empty) / placeholder ("Na", 0, ...) / substantive
        NA_TXT = {"na", "n/a", "nil", "none", "no", "0", "-", "nan"}
        def row_kind(r):
            vals = [r[v] for v in datavars if pd.notna(r[v])]
            if not vals:
                return "blank"
            if all((isinstance(x, str) and x.strip().lower() in NA_TXT) or
                   (isinstance(x, (int, float)) and not isinstance(x, bool) and x == 0) for x in vals):
                return "placeholder"
            return "substantive"
        long["__kind"] = long.apply(row_kind, axis=1)
        kinds = long["__kind"].value_counts()
        if kinds.get("blank", 0) or kinds.get("placeholder", 0):
            checks.append((f"{name}: repeat rows with no real data (kept as recorded)",
                           int(kinds.get("blank", 0) + kinds.get("placeholder", 0)),
                           f"{int(kinds.get('blank', 0))} fully blank, {int(kinds.get('placeholder', 0))} placeholder entries "
                           f"('Na' / 0) — the form opens one item by default, so these mostly sit in item 1 where {count_var} = 0 or missing"))
        # reported count vs substantive items recorded
        n_items = long[long["__kind"] == "substantive"].groupby("_parent_index").size()
        mi = main.set_index("_index")
        cnt = mi[count_var]
        rec = n_items.reindex(cnt.index).fillna(0).astype(int)
        bad = cnt.index[cnt.fillna(0) != rec]
        if len(bad):
            det = ", ".join(f"{mi.at[i, 'farmers_id']} ({'NA' if pd.isna(cnt[i]) else int(cnt[i])} vs {rec[i]})" for i in bad)
            checks.append((f"{name}: {count_var} ≠ items with real data", len(bad),
                           "farmers_id (count reported vs items recorded): " + det))
        print(f"{name:8s} rows={len(long):4d} parents={long['_parent_index'].nunique():3d} max_items={max_item} new_cols={len(cols)}")

    # ---- build merged frame: insert each block right after its count variable ----
    merged = main.set_index("_index", drop=False)
    new_cols_all = [c for w in wide_blocks.values() for c in w.columns]
    clash = set(new_cols_all) & set(main.columns)
    assert not clash, f"column name clash: {clash}"
    assert len(new_cols_all) == len(set(new_cols_all)), "duplicate new column names"

    order = []
    for c in main.columns:
        order.append(c)
        if c in wide_blocks:
            order.extend(wide_blocks[c].columns)
    merged = merged.join(pd.concat(wide_blocks.values(), axis=1), how="left")
    merged = merged[order].reset_index(drop=True).copy()
    assert len(merged) == n_main

    # farmer id: numbers as numbers, keep any text id (e.g. "502 new") as text
    fid_raw = merged["farmers_id"].astype(str).str.strip()
    merged["farmers_id"] = [int(x) if re.fullmatch(r"\d+", x) else x for x in fid_raw]

    # ---- payment ----
    pay = read_payment(PAY)
    pay.columns = [c.strip() for c in pay.columns]
    pay["farmer_id"] = pay["farmer_id"].astype(str).str.strip()
    pay["Payment"] = pay["Payment"].astype(str).str.strip()
    assert set(pay["Payment"]) <= {"Yes", "No"}
    assert pay["farmer_id"].is_unique
    pmap = dict(zip(pay["farmer_id"], pay["Payment"]))
    in_list = fid_raw.isin(pmap.keys())
    payment = fid_raw.map(lambda x: 1 if pmap.get(x) == "Yes" else 0)
    pos = list(merged.columns).index("farmers_id") + 1
    merged.insert(pos, "round", 1)
    merged.insert(pos + 1, "payment", payment.astype(int))
    merged.insert(pos + 2, "in_payment_list", in_list.astype(int))

    # ---- Section 3: activity costs, per-acre costs, outliers ----
    cst, cost_review, cost_labels = C.compute(merged)
    cost_summary, cost_outl, cost_flags = C.outliers(merged, cst, cost_labels)
    cst["cpa_total_outlier"] = cost_flags["out_total"]
    cst["cost_outlier_n"] = cost_flags[[f"out_{k}" for k in C.ORDER]].sum(axis=1)
    anchor = list(merged.columns).index("f_other_paydate_4") + 1
    merged = pd.concat([merged.iloc[:, :anchor], cst, merged.iloc[:, anchor:]], axis=1)
    cost_cols = list(cst.columns)

    def cost_label(c):
        if c in ("cost_total", "cpa_total"):
            base = "TOTAL cost of rice cultivation, Virippu 2025 (sum of Sections A–F)"
        elif c[4:7] == "sec" or c[5:8] == "sec":
            s = c[-1]
            base = f"Section {s} total — {C.SECTIONS[s]} (sum of its activities)"
        elif c.startswith(("cost_", "cpa_")) and c.split("_", 1)[1] in cost_labels:
            k = c.split("_", 1)[1]
            lab, sec, how = cost_labels[k]
            base = f"{sec}. {lab} = {how}"
        else:
            base = None
        if c == "cost_n_missing":
            return "Number of activities with at least one cost item that could not be computed (999 code or quantity/price half-filled); such items count as 0 in totals"
        if c == "cpa_total_outlier":
            return "1 = total cost per acre is an outlier (|modified z| > 3.5 on log scale); see sheet cost_outliers"
        if c == "cost_outlier_n":
            return "Number of activities (of 27) whose cost per acre is flagged as an outlier for this household"
        if c.startswith("cost_"):
            return f"Cost (Rs.): {base}"
        return f"Cost per acre (Rs./acre): {base}, ÷ rice_acres_2025 (missing if rice area is 0)"

    # ---- data checks ----
    dups = merged[merged["farmers_id"].duplicated(keep=False)]
    checks.insert(0, ("Duplicate farmers_id remaining", dups["farmers_id"].nunique(),
                      "; ".join(f"{f}: _index {', '.join(map(str, g['_index']))}" for f, g in dups.groupby("farmers_id")) or "none"))
    checks.insert(0, ("farmers_id corrected (user request; payment recomputed on new ID)", len(fix_notes),
                      "; ".join(n + (f" — payment list: {pmap[n.split('→ ')[1]]}" if n.split('→ ')[1] in pmap else " — not in payment list")
                                for n in fix_notes)))
    txt = [x for x in fid_raw if not re.fullmatch(r"\d+", x)]
    checks.insert(1, ("Non-numeric farmers_id (kept as text)", len(txt), ", ".join(txt)))
    missing = sorted(set(pmap) - set(fid_raw), key=lambda x: int(x) if x.isdigit() else 10**9)
    checks.insert(2, ("Payment-list IDs not found in survey", len(missing),
                      ", ".join(f"{m} ({pmap[m]})" for m in missing)))
    checks.insert(3, ("Survey rows coded payment = 1 / 0", f"{int(payment.sum())} / {int((payment == 0).sum())}",
                      f"of which in payment list with 'No': {int((in_list & (payment == 0)).sum())}; not in list: {int((~in_list).sum())}"))
    empty = [c for c in main.columns if main[c].isna().all()]
    checks.append(("Columns from main sheet that are completely empty (kept)", len(empty), ", ".join(empty)))
    acres = pd.to_numeric(merged["rice_acres_2025"], errors="coerce")
    z_ac = merged.loc[acres == 0, "farmers_id"].astype(str).tolist()
    checks.append(("Section 3 costs: rice_acres_2025 = 0 although costs are reported → per-acre costs missing", len(z_ac), ", ".join(z_ac)))
    big = merged.loc[acres > 20, ["farmers_id"]].assign(a=acres[acres > 20])
    checks.append(("Section 3 costs: rice_acres_2025 > 20 acres (check — costs reported are small for this area)", len(big),
                   ", ".join(f"{r.farmers_id} ({r.a:g} acres)" for r in big.itertuples())))
    checks.append(("Section 3 costs: 1st-dose fertiliser recorded differently by form version", int((main["__version__"] == "vhWzpqVLG8rKBa5kqM2f3d").sum()),
                   "Older form: one fertiliser item in the main sheet (c_fert1_name / _qty / _unit / _price); newer form: up to 2 items in a repeat group "
                   "(c_fert1_*_1, c_fert1_*_2). cost_c_fert1 uses whichever the household has."))
    checks.append(("Section 3 costs: households with total cost per acre flagged as outlier", int(cst["cpa_total_outlier"].sum()),
                   ", ".join(merged.loc[cst["cpa_total_outlier"] == 1, "farmers_id"].astype(str))))
    checks.insert(0, ("Form versions combined", 2,
                      "; ".join(f"{v}: {n} interviews" for v, n in main["__version__"].value_counts().items())))

    # ---- codebook ----
    cb = []
    for i, c in enumerate(merged.columns, 1):
        if c in codebook_extra:
            src, sec, k, base = codebook_extra[c]
            cb.append([i, c, src, sec, k, f"{label_for(base)} [item {k}]".strip()])
        elif c in cost_cols:
            cb.append([i, c, "Computed from Section 3", "Sec 3: cost of rice cultivation, Virippu 2025", None, cost_label(c)])
        elif c == "round":
            cb.append([i, c, "Added", "Survey round", None, "Survey round (1 for all rows in this file)"])
        elif c == "payment":
            cb.append([i, c, "RCT_payment_round1.xls", "RCT payment", None,
                       "1 = Payment 'Yes' in RCT_payment_round1; 0 = 'No' or farmer not in the payment list"])
        elif c == "in_payment_list":
            cb.append([i, c, "RCT_payment_round1.xls", "RCT payment", None,
                       "1 = farmer appears in RCT_payment_round1 (Yes or No); 0 = not listed"])
        else:
            cb.append([i, c, MAIN_SHEET, "Main sheet", None, label_for(c)])
    codebook = pd.DataFrame(cb, columns=["col_no", "variable", "source_sheet", "section", "item_no", "label"])
    checks_df = pd.DataFrame(checks, columns=["check", "count", "details"])

    with pd.ExcelWriter(OUT, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm", date_format="yyyy-mm-dd") as xw:
        merged.to_excel(xw, sheet_name="data", index=False)
        codebook.to_excel(xw, sheet_name="codebook", index=False)
        checks_df.to_excel(xw, sheet_name="data_checks", index=False)
        cost_summary.to_excel(xw, sheet_name="cost_summary", index=False, startrow=2)
        cost_outl.sort_values(["direction", "variable", "times the median"], ascending=[True, True, False]) \
            .to_excel(xw, sheet_name="cost_outliers", index=False, startrow=2)
        cost_review.to_excel(xw, sheet_name="cost_review", index=False, startrow=2)

    # ---- formatting ----
    date_only = {c for c in merged.columns
                 if pd.api.types.is_datetime64_any_dtype(merged[c])
                 and (merged[c].dropna().dt.normalize() == merged[c].dropna()).all()}
    wb = load_workbook(OUT)
    base_font = Font(name="Arial", size=10)
    head_font = Font(name="Arial", size=10, bold=True)
    head_fill = PatternFill("solid", start_color="DDE4EE")
    new_fill = PatternFill("solid", start_color="E8F1E4")     # header shade for columns added from repeat sheets / payment
    cost_fill = PatternFill("solid", start_color="FCE9D9")    # header shade for computed cost variables
    added = set(codebook_extra) | {"round", "payment", "in_payment_list"}
    COST_SHEETS = {
        "cost_summary": ("Cost per acre (Rs./acre) by activity, Virippu 2025 — distribution and outlier counts",
                         "Statistics over households with cost > 0. Outlier = |modified z-score| > 3.5 on log(cost per acre) "
                         "(Iglewicz & Hoaglin), i.e. outside the lower/upper bounds shown. The Q3 + 3×IQR count is a second, raw-scale check."),
        "cost_outliers": ("Flagged per-acre costs, one row per household × activity",
                          "‘likely cause’ compares the entry with the median quantity and price for that activity. Values are as recorded — nothing has been corrected."),
        "cost_review": ("Decisions taken while computing costs (999 codes, half-filled entries, 'amount paid' read as a unit price)",
                        "Other-item questions ask the total amount paid; where amount ÷ quantity is implausibly low the amount is read as a unit price. "
                        "Ambiguous cases are kept as totals and listed here to check."),
    }
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = base_font
        hdr_row = 3 if ws.title in COST_SHEETS else 1
        for cell in ws[hdr_row]:
            cell.font = head_font
            cell.fill = (cost_fill if (ws.title == "data" and cell.value in cost_cols) else
                         new_fill if (ws.title == "data" and cell.value in added) else head_fill)
            cell.alignment = Alignment(vertical="center", wrap_text=ws.title in COST_SHEETS)
        ws.freeze_panes = "C2" if ws.title == "data" else ("A4" if ws.title in COST_SHEETS else "A2")
        if ws.title in COST_SHEETS:
            ws["A1"], ws["A2"] = COST_SHEETS[ws.title]
            ws["A1"].font = Font(name="Arial", size=12, bold=True)
            ws["A2"].font = Font(name="Arial", size=10, italic=True)
            ws.auto_filter.ref = f"A3:{get_column_letter(ws.max_column)}{ws.max_row}"
            ws.row_dimensions[3].height = 30
        else:
            ws.auto_filter.ref = ws.dimensions
    for name, widths in {"cost_summary": [8, 46, 18] + [11] * 14,
                         "cost_outliers": [10, 8, 8, 40, 18, 9, 10, 14, 14, 12, 10, 10, 110],
                         "cost_review": [10, 8, 11, 32, 26, 10, 18, 14, 14, 100]}.items():
        ws = wb[name]
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00" if ws.cell(3, cell.column).value in ("modified z (log)", "times the median", "rice_acres_2025") else "#,##0"
    ws = wb["data"]
    for j, c in enumerate(merged.columns, 1):
        if c in cost_cols and c.startswith(("cost_", "cpa_")) and c not in ("cost_n_missing", "cpa_total_outlier", "cost_outlier_n"):
            for (cell,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                cell.number_format = "#,##0"
    ws = wb["data"]
    for j, c in enumerate(merged.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(len(str(c)) + 2, 10), 40)
        if c in date_only:
            for (cell,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                cell.number_format = "yyyy-mm-dd"
    cbs = wb["codebook"]
    for col, wdt in zip("ABCDEF", [8, 30, 48, 42, 9, 100]):
        cbs.column_dimensions[col].width = wdt
    dc = wb["data_checks"]
    for col, wdt in zip("ABC", [58, 12, 120]):
        dc.column_dimensions[col].width = wdt
    for (cell,) in dc.iter_rows(min_row=2, min_col=3, max_col=3):
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT)
    print("final shape:", merged.shape, "added columns:", len(added))
    print(checks_df.to_string())


# =====================================================================================
# STEP 2 — baseline balance by payment  ->  Balance_table_payment_round1.xlsx
# =====================================================================================
def balance_table(DATA, OUT):
    d = pd.read_excel(DATA, sheet_name="data")
    T = d["payment"] == 1

    def num(c):
        return pd.to_numeric(d[c], errors="coerce")

    # ---------------- analysis-level cleaning (the dataset itself is not changed) ----------------
    recode_log = []
    def clean(c, s):
        s = s.copy()
        codes = s.isin([999, 9999])
        if codes.any():
            recode_log.append((c, int(codes.sum()), "999 / 9999 treated as 'don't know' code → missing"))
            s[codes] = np.nan
        return s

    V = {}   # variable name -> (label, series, type)  type: 'bin' | 'cont' | 'money'
    def add(key, label, s, kind):
        V[key] = (label, s.astype(float), kind)

    # Panel A: respondent
    add("female", "Female respondent (share)", (num("gender") == 1), "bin")
    age = num("age")
    bad_age = age < 18
    recode_log.append(("age", int(bad_age.sum()), f"age < 18 for a respondent (values {sorted(age[bad_age].astype(int).tolist())}) → missing"))
    add("age", "Age (years)", age.where(~bad_age), "cont")
    add("hh_head", "Respondent is household head (share)", num("rel_head") == 1, "bin")
    edu = num("education")
    add("edu_low", "Education: not literate / below primary (share)", edu.isin([1, 2, 3]), "bin")
    add("edu_primmid", "Education: primary or middle (share)", edu.isin([4, 5]), "bin")
    add("edu_sec", "Education: secondary (share)", edu == 6, "bin")
    add("edu_hsplus", "Education: higher secondary or above (share)", edu.isin([10, 11, 12, 13]), "bin")
    add("emp_selfag", "Self-employed in agriculture (share)", num("employment") == 1, "bin")
    add("ag_sole", "Sole decision-maker: farm investment (share)", num("ag_decision") == 1, "bin")
    add("hh_sole", "Sole decision-maker: household investment (share)", num("hh_decision") == 1, "bin")

    # Panel B: household demographics
    hh = num("hh_members")
    bad_hh = hh > 20
    recode_log.append(("hh_members", int(bad_hh.sum()),
                       f"household size > 20 (values {sorted(hh[bad_hh].astype(int).tolist())}; inconsistent with working-age + children counts) → missing"))
    add("hh_members", "Household size (members)", hh.where(~bad_hh), "cont")
    add("working_age", "Working-age members (18–60)", num("working_age"), "cont")
    add("children", "Children (< 18)", num("children"), "cont")
    add("migrants", "Migrant members", num("migrants"), "cont")
    add("offfarm_count", "Members in off-farm / non-farm work", num("offfarm_count"), "cont")

    # Panel C: farm
    add("total_area", "Total cultivated area (acres)", num("total_area"), "cont")
    add("land_owned", "Land owned (acres)", num("land_owned"), "cont")
    add("lease_in", "Leases in land (share)", (num("land_leased_in") > 0).where(num("land_leased_in").notna()), "bin")
    add("land_leased_in", "Land leased in (acres)", num("land_leased_in"), "cont")
    add("rice_acres_2025", "Rice area, Virippu 2025 (acres)", num("rice_acres_2025"), "cont")
    add("rice_prod_2025", "Rice yield, Virippu 2025 (kg/acre)", num("rice_prod_2025"), "cont")
    add("rice_acres_2026", "Planned rice area, Virippu 2026 (acres)", num("rice_acres_2026"), "cont")
    add("other_crops", "Grows crops other than rice (share)", num("other_crops_count") > 0, "bin")

    # Panel D: consumption expenditure (last 30 days)
    exp_items = [("exp_food_amt", "Food"), ("exp_util_amt", "Housing & utilities"), ("exp_comm_amt", "Transport & communication"),
                 ("exp_health_amt", "Health"), ("exp_edu_amt", "Education"), ("exp_personal_amt", "Personal & household goods"),
                 ("exp_misc_amt", "Miscellaneous")]
    exp_clean = {}
    for c, lab in exp_items:
        exp_clean[c] = clean(c, num(c))
        add(c, f"{lab} (Rs., last 30 days)", exp_clean[c], "money")
    tot = pd.concat(exp_clean.values(), axis=1)
    add("exp_sum", "Total consumption, sum of 7 items (Rs., last 30 days)", tot.sum(axis=1, min_count=7), "money")

    # Panel E: income and credit (2025)
    inc_items = [("inc_rice_vir_amt", "Income from rice, Virippu 2025"), ("inc_rice_mun_amt", "Income from rice, Mundakan 2025-26"),
                 ("inc_agnet_amt", "Total agricultural income"), ("inc_offfarm_amt", "Off-farm income"),
                 ("inc_salary_amt", "Salaried wages"), ("inc_nonfarm_amt", "Non-farm (non-salaried) income"),
                 ("inc_pension_amt", "Pension"), ("inc_other_amt", "Other income"), ("inc_loan_amt", "Loan outstanding")]
    for c, lab in inc_items:
        add(c, f"{lab} (Rs.)", clean(c, num(c)), "money")
    loan = V["inc_loan_amt"][1]
    add("has_loan", "Has an outstanding loan (share)", (loan > 0).where(loan.notna()), "bin")

    # Panel F: PM-KISAN, welfare, banking
    add("pmkisan", "Ever received PM-KISAN (share)", num("pmkisan_received") == 1, "bin")
    add("smartphone", "Owns a smartphone (share)", num("has_phone/2") == 1, "bin")
    add("sms_alert", "Bank SMS alerts enabled (share)", num("sms_alert") == 1, "bin")
    add("has_upi", "Has UPI / digital wallet (share)", num("has_upi") == 1, "bin")
    add("welfare_count", "Other welfare schemes received (number)", num("welfare_count"), "cont")
    add("bank_distance", "Distance to bank (km)", clean("bank_distance", num("bank_distance")), "cont")
    add("bank_visits", "Bank visits per year", clean("bank_visits", num("bank_visits")), "cont")
    add("upi_dealers", "Pays input dealers by UPI (share)", num("pay_dealers/3") == 1, "bin")
    add("upi_labour", "Pays labourers by UPI (share)", num("pay_labourers/3") == 1, "bin")
    add("upi_grocery", "Pays grocery shops by UPI (share)", num("pay_grocery/3") == 1, "bin")

    PANELS = [
        ("A. Respondent characteristics", ["female", "age", "hh_head", "edu_low", "edu_primmid", "edu_sec", "edu_hsplus", "emp_selfag", "ag_sole", "hh_sole"]),
        ("B. Household demographics", ["hh_members", "working_age", "children", "migrants", "offfarm_count"]),
        ("C. Farm characteristics", ["total_area", "land_owned", "lease_in", "land_leased_in", "rice_acres_2025", "rice_prod_2025", "rice_acres_2026", "other_crops"]),
        ("D. Household consumption expenditure (last 30 days)", [c for c, _ in exp_items] + ["exp_sum"]),
        ("E. Income and credit (2025)", [c for c, _ in inc_items] + ["has_loan"]),
        ("F. PM-KISAN, welfare schemes and banking", ["pmkisan", "smartphone", "sms_alert", "has_upi", "welfare_count", "bank_distance", "bank_visits", "upi_dealers", "upi_labour", "upi_grocery"]),
    ]
    assert sorted(k for _, ks in PANELS for k in ks) == sorted(V), "panel list mismatch"

    def stars(p):
        return "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.10 else ""

    rows = []
    for panel, keys in PANELS:
        rows.append({"Variable": panel, "_panel": True})
        for k in keys:
            lab, s, kind = V[k]
            t, c = s[T].dropna(), s[~T].dropna()
            diff = t.mean() - c.mean()
            tt = stats.ttest_ind(t, c, equal_var=False)          # Welch two-sample t-test
            se = np.sqrt(t.var(ddof=1) / len(t) + c.var(ddof=1) / len(c))
            nd = diff / np.sqrt((t.var(ddof=1) + c.var(ddof=1)) / 2) if (t.var() + c.var()) > 0 else np.nan
            mw = stats.mannwhitneyu(t, c, alternative="two-sided").pvalue
            rows.append({"Variable": lab, "key": k, "kind": kind,
                         "N_T": len(t), "Mean_T": t.mean(), "SD_T": t.std(ddof=1),
                         "N_C": len(c), "Mean_C": c.mean(), "SD_C": c.std(ddof=1),
                         "Diff": diff, "SE": se, "Stars": stars(tt.pvalue), "p_t": tt.pvalue,
                         "NormDiff": nd, "p_mw": mw, "_panel": False})
    bal = pd.DataFrame(rows)

    # ---------------- categorical distributions: chi-square tests ----------------
    cat_specs = [("education", "Education (10 categories)"), ("employment", "Employment type"),
                 ("rel_head", "Relationship to household head"), ("ag_decision", "Role in farm investment decisions"),
                 ("hh_decision", "Role in household investment decisions")]
    cat_rows = []
    for c, lab in cat_specs:
        tab = pd.crosstab(num(c), d["payment"])
        chi2, p, dof, exp = stats.chi2_contingency(tab)
        cat_rows.append([lab, c, int(tab.values.sum()), round(chi2, 3), int(dof), p, int((exp < 5).sum()), int(exp.size)])
    cat = pd.DataFrame(cat_rows, columns=["Characteristic", "variable", "N", "Chi-square", "df", "p-value", "cells with expected < 5", "cells"])

    # ---------------- joint orthogonality test (linear probability model, HC1 robust) ----------------
    joint_keys = ["female", "age", "hh_head", "edu_sec", "edu_hsplus", "emp_selfag", "ag_sole", "hh_members", "working_age",
                  "children", "migrants", "offfarm_count", "total_area", "land_owned", "rice_acres_2025", "rice_prod_2025",
                  "other_crops", "exp_sum", "inc_rice_vir_amt", "inc_rice_mun_amt", "inc_loan_amt", "pmkisan", "smartphone",
                  "sms_alert", "has_upi", "welfare_count", "bank_distance", "bank_visits"]
    X = pd.concat({k: V[k][1] for k in joint_keys}, axis=1)
    y = d["payment"].astype(float)
    ok = X.notna().all(axis=1)
    Xc = np.column_stack([np.ones(ok.sum()), X[ok].values])
    yc = y[ok].values
    n, kx = Xc.shape
    q = kx - 1
    XtX_inv = np.linalg.inv(Xc.T @ Xc)
    H = XtX_inv @ Xc.T                                    # beta = H @ y
    lev = np.einsum("ij,ji->i", Xc, H)                    # leverage h_ii

    def joint_stats(yv):
        beta = H @ yv
        e = yv - Xc @ beta
        meat1 = (Xc.T * e**2) @ Xc
        meat3 = (Xc.T * (e / (1 - lev))**2) @ Xc
        V1 = n / (n - kx) * XtX_inv @ meat1 @ XtX_inv
        V3 = XtX_inv @ meat3 @ XtX_inv
        b = beta[1:]
        F1 = float(b @ np.linalg.solve(V1[1:, 1:], b)) / q
        F3 = float(b @ np.linalg.solve(V3[1:, 1:], b)) / q
        ssr_u = float(e @ e); ssr_r = float(((yv - yv.mean()) ** 2).sum())
        Fo = ((ssr_r - ssr_u) / q) / (ssr_u / (n - kx))
        return beta, V1, F1, F3, Fo, 1 - ssr_u / ssr_r

    beta, V_hc1, F_rob, F_hc3, F_ols, r2 = joint_stats(yc)
    p_rob = 1 - stats.f.cdf(F_rob, q, n - kx)
    p_hc3 = 1 - stats.f.cdf(F_hc3, q, n - kx)
    p_ols = 1 - stats.f.cdf(F_ols, q, n - kx)
    # randomization inference: re-shuffle payment across households, recompute the F statistics
    rng = np.random.default_rng(20261026)
    NPERM = 5000
    perm = np.array([joint_stats(rng.permutation(yc))[2:5] for _ in range(NPERM)])
    ri_rob = (np.sum(perm[:, 0] >= F_rob) + 1) / (NPERM + 1)
    ri_hc3 = (np.sum(perm[:, 1] >= F_hc3) + 1) / (NPERM + 1)
    ri_ols = (np.sum(perm[:, 2] >= F_ols) + 1) / (NPERM + 1)
    se_b = np.sqrt(np.diag(V_hc1))
    coef = pd.DataFrame({"Variable": ["Constant"] + [V[k][0] for k in joint_keys],
                         "Coefficient": beta, "Robust SE (HC1)": se_b,
                         "t": beta / se_b, "p-value": 2 * (1 - stats.t.cdf(np.abs(beta / se_b), n - kx))})

    # ---------------- write workbook ----------------
    nT, nC = int(T.sum()), int((~T).sum())
    hdr = ["Variable", "N", "Mean", "SD", "N", "Mean", "SD", "Difference (T − C)", "", "SE", "p-value (t-test)", "Normalized difference", "p-value (Mann–Whitney)"]
    wb_rows = []
    for _, r in bal.iterrows():
        if r["_panel"]:
            wb_rows.append([r["Variable"]] + [None] * 12)
        else:
            wb_rows.append([r["Variable"], r["N_T"], r["Mean_T"], r["SD_T"], r["N_C"], r["Mean_C"], r["SD_C"],
                            r["Diff"], r["Stars"], r["SE"], r["p_t"], r["NormDiff"], r["p_mw"]])
    n_sig10 = int((bal["p_t"] < 0.10).sum()); n_sig5 = int((bal["p_t"] < 0.05).sum()); n_sig1 = int((bal["p_t"] < 0.01).sum())
    n_vars = int((~bal["_panel"]).sum())
    n_nd = int((bal["NormDiff"].abs() > 0.25).sum())

    notes = [
        ["Treatment definition", f"payment = 1 (Payment = 'Yes' in RCT_payment_round1): {nT} households. Control: payment = 0: {nC} households "
                                 f"(5 listed with Payment = 'No' + {nC - 5} not in the payment list)."],
        ["Data", "Survey Round 1 merged dataset (Survey_Round_1_merged_wide.xlsx, sheet 'data'), 322 households, after farmer ID corrections (165→166, 201→200)."],
        ["Difference and p-value", "Difference = mean(treatment) − mean(control). p-value from a two-sample Welch t-test (unequal variances). "
                                  "Shares are means of 0/1 indicators, so the same test is a difference-in-proportions test. *** p<0.01, ** p<0.05, * p<0.10."],
        ["Normalized difference", "Difference ÷ √[(SD²_T + SD²_C)/2] (Imbens & Rubin 2015). |value| > 0.25 is the usual rule of thumb for a meaningful imbalance."],
        ["Mann–Whitney", "Non-parametric rank-sum test, reported because incomes, expenditures and land are highly skewed with large outliers."],
        ["Joint test", f"OLS of payment on {q} baseline covariates (linear probability model), F-test that all slopes are zero: conventional, HC3- and HC1-robust, "
                       f"each with an F-distribution p-value and a randomization-inference p-value ({NPERM:,} random re-assignments of payment). "
                       f"Complete cases only: N = {n}. The randomization-inference p-values are the most reliable here."],
        ["Missing values", "Each row uses all households with a non-missing value; N per group is shown. Skipped questions are left missing (not set to zero)."],
    ]
    for c, cnt, why in recode_log:
        if cnt:
            notes.append([f"Recode: {c}", f"{cnt} value(s): {why}. Applied for this analysis only; the merged dataset is unchanged."])
    notes.append(["Sensitivity to 999 recode", "Re-running the t-tests with the 999 values kept as real amounts does not change which variables are significant at 10%."])
    notes.append(["Expenditure total", "The reported total (exp_total) does not equal the sum of the 7 items for most households, so the table uses the sum of the items."])

    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        pd.DataFrame(wb_rows, columns=hdr).to_excel(xw, sheet_name="balance_table", index=False, startrow=3)
        cat.to_excel(xw, sheet_name="categorical_tests", index=False, startrow=2)
        coef.to_excel(xw, sheet_name="joint_test", index=False, startrow=8)
        pd.DataFrame(notes, columns=["Item", "Note"]).to_excel(xw, sheet_name="notes", index=False)

    wb = load_workbook(OUT)
    F = lambda **k: Font(**{"name": "Arial", "size": 10, **k})
    thin = Side(style="thin", color="000000")
    hfill = PatternFill("solid", start_color="DDE4EE")
    pfill = PatternFill("solid", start_color="F2F2F2")

    # balance table
    ws = wb["balance_table"]
    for row in ws.iter_rows():
        for cell in row:
            cell.font = F()
    ws["A1"] = "Table: Baseline balance of household socio-economic characteristics by payment status"
    ws["A1"].font = F(bold=True, size=12)
    ws["A2"] = (f"Treatment = payment received (N = {nT}); Control = no payment (N = {nC}). "
                f"Welch t-test p-values; *** p<0.01, ** p<0.05, * p<0.10. See 'notes' for definitions and recodes.")
    ws["A2"].font = F(italic=True)
    ws["B3"] = f"Treatment (payment = 1)"; ws.merge_cells("B3:D3")
    ws["E3"] = f"Control (payment = 0)"; ws.merge_cells("E3:G3")
    ws["H3"] = "Difference"; ws.merge_cells("H3:M3")
    ws.merge_cells("H4:I4")
    for cell in list(ws[3]) + list(ws[4]):
        cell.font = F(bold=True); cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[4].height = 42
    for c in range(1, 14):
        ws.cell(row=3, column=c).border = Border(top=thin)
        ws.cell(row=4, column=c).border = Border(bottom=thin)
    r0 = 5
    for i, (_, r) in enumerate(bal.iterrows()):
        rr = r0 + i
        if r["_panel"]:
            ws.cell(rr, 1).font = F(bold=True)
            for c in range(1, 14):
                ws.cell(rr, c).fill = pfill
            continue
        kind = r["kind"]
        fmt = "0.000" if kind == "bin" else "#,##0" if kind == "money" else "0.00"
        for c in (3, 4, 6, 7, 8, 10):
            ws.cell(rr, c).number_format = fmt
        for c in (2, 5):
            ws.cell(rr, c).number_format = "0"
        for c in (11, 12, 13):
            ws.cell(rr, c).number_format = "0.000"
        ws.cell(rr, 1).alignment = Alignment(indent=1)
        if r["p_t"] < 0.10:
            for c in (8, 9, 11):
                ws.cell(rr, c).font = F(bold=True)
        if abs(r["NormDiff"]) > 0.25:
            ws.cell(rr, 12).font = F(bold=True)
    last = r0 + len(bal) - 1
    for c in range(1, 14):
        ws.cell(last, c).border = Border(bottom=thin)
    summ = last + 2
    ws.cell(summ, 1, f"Joint test of {q} covariates (N = {n}): conventional F({q}, {n - kx}) = {F_ols:.2f}, randomization-inference p = {ri_ols:.3f}; "
                     f"robust HC3 F = {F_hc3:.2f}, randomization-inference p = {ri_hc3:.3f}  (details in 'joint_test')").font = F(bold=True)
    ws.cell(summ + 1, 1, f"Variables with p < 0.10: {n_sig10} of {n_vars}  (p < 0.05: {n_sig5}; p < 0.01: {n_sig1}).  "
                         f"Expected by chance at 10%: about {0.1 * n_vars:.0f}.  |Normalized difference| > 0.25: {n_nd}.")
    ws.cell(summ + 2, 1, "Bold = difference significant at 10% or |normalized difference| > 0.25.").font = F(italic=True)
    widths = [56, 6, 11, 11, 6, 11, 11, 13, 5, 11, 10, 12, 13]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:4"

    # categorical tests
    ws = wb["categorical_tests"]
    for row in ws.iter_rows():
        for cell in row:
            cell.font = F()
    ws["A1"] = "Chi-square tests of independence: distribution of categorical characteristics by payment status"
    ws["A1"].font = F(bold=True, size=12)
    for cell in ws[3]:
        cell.font = F(bold=True); cell.fill = hfill; cell.alignment = Alignment(wrap_text=True, vertical="center")
    for (cell,) in ws.iter_rows(min_row=4, min_col=6, max_col=6):
        cell.number_format = "0.000"
    r = 4 + len(cat) + 1
    ws.cell(r, 1, "Many cells have expected counts below 5, so treat these chi-square p-values as approximate; the grouped indicators in the balance table are the main test.").font = F(italic=True)
    for j, w in enumerate([40, 14, 7, 11, 5, 9, 14, 7], 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # joint test
    ws = wb["joint_test"]
    for row in ws.iter_rows():
        for cell in row:
            cell.font = F()
    ws["A1"] = "Joint orthogonality test: OLS of payment (1/0) on baseline covariates"
    ws["A1"].font = F(bold=True, size=12)
    ws["A2"] = f"Observations (complete cases): {n}    Covariates (excl. constant): {q}    R-squared: {r2:.3f}    Permutations: {NPERM:,}"
    ws["A3"] = "Test that all slopes = 0"; ws["B3"] = "F statistic"; ws["C3"] = "p-value (F distribution)"; ws["D3"] = "p-value (randomization inference)"
    for cell in ws[3]:
        cell.font = F(bold=True); cell.fill = hfill; cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, (k, f_, p_, ri) in enumerate([("Conventional (homoskedastic) F-test", F_ols, p_ols, ri_ols),
                                         ("Robust F-test, HC3", F_hc3, p_hc3, ri_hc3),
                                         ("Robust F-test, HC1", F_rob, p_rob, ri_rob)], 4):
        ws.cell(i, 1, k); ws.cell(i, 2, round(f_, 3)); ws.cell(i, 3, round(p_, 4)); ws.cell(i, 4, round(ri, 4))
    ws["A7"] = ("Randomization inference re-assigns payment at random across the same households and asks how often the F statistic is at least as large "
                "as observed. It does not rely on the F distribution, which robust tests with many covariates in a small sample tend to over-reject.")
    ws["A7"].font = F(italic=True); ws["A7"].alignment = Alignment(wrap_text=True, vertical="top"); ws.merge_cells("A7:E7"); ws.row_dimensions[7].height = 42
    for cell in ws[9]:
        cell.font = F(bold=True); cell.fill = hfill
    for row in ws.iter_rows(min_row=10, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = "0.0000"
    ws.column_dimensions["A"].width = 56
    for col in "BCDE":
        ws.column_dimensions[col].width = 16

    # notes
    ws = wb["notes"]
    for row in ws.iter_rows():
        for cell in row:
            cell.font = F(); cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = F(bold=True); cell.fill = hfill
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    wb.save(OUT)

    # console summary
    pd.set_option("display.width", 250)
    show = bal[~bal["_panel"]][["Variable", "N_T", "Mean_T", "N_C", "Mean_C", "Diff", "p_t", "NormDiff", "p_mw"]]
    print(show.round(3).to_string(index=False))
    print(f"\nJoint: OLS F={F_ols:.3f} p={p_ols:.4f} ri={ri_ols:.4f} | HC3 F={F_hc3:.3f} p={p_hc3:.4f} ri={ri_hc3:.4f} | "
          f"HC1 F={F_rob:.3f} p={p_rob:.4f} ri={ri_rob:.4f} | N={n}, q={q}, R2={r2:.3f}")
    print(f"sig10={n_sig10}, sig5={n_sig5}, sig1={n_sig1} of {n_vars}; |nd|>0.25: {n_nd}")
    print(cat.round(3).to_string(index=False))
    print(recode_log)


# =====================================================================================
# STEP 3 — per-acre cost components by payment  ->  Cost_comparison_payment_round1.xlsx
# =====================================================================================
def cost_comparison(DATA, OUT):
    d = pd.read_excel(DATA, sheet_name="data").copy()
    outl = pd.read_excel(DATA, sheet_name="cost_outliers", header=2)
    T = d["payment"] == 1
    old = d["__version__"] == "vhWzpqVLG8rKBa5kqM2f3d"        # older form version (1st-dose fertiliser recorded as one item)
    acres = pd.to_numeric(d["rice_acres_2025"], errors="coerce")

    # totals are complete in both form versions; the *_x names are kept only for the code below
    d["cost_secC_x"] = d["cost_secC"]
    d["cost_total_x"] = d["cost_total"]
    for c in ("secC_x", "total_x"):
        d[f"cpa_{c}"] = (d[f"cost_{c}"] / acres).where(acres > 0)

    # outlier mask: (row, variable) pairs flagged in the cost_outliers sheet; derived totals get their own check
    flag = pd.DataFrame(False, index=d.index, columns=[c for c in d.columns if c.startswith("cpa_")])
    idx_of = dict(zip(d["_index"], d.index))
    for r in outl.itertuples():
        flag.at[idx_of[r._2], r.variable] = True
    flag["cpa_secC_x"] = flag["cpa_secC"]
    flag["cpa_total_x"] = flag["cpa_total"]

    LAB = {k: v for k, v in [(k, l) for k, l, *_ in C.FIXED] + [(k, l) for k, l, *_ in C.FERT] + [(k, l) for k, l, *_ in C.OTHER]}
    PANELS = [(f"{s}. {C.SECTIONS[s]}", [k for k in C.ORDER if k[0] == s.lower()]) for s in C.SECTIONS]
    TOTALS = [("secA", "Section A: Pre-nursery preparation"), ("secB", "Section B: Nursery preparation"),
              ("secC", "Section C: Transplanting"),
              ("secD", "Section D: Weedicide / fertiliser application"), ("secE", "Section E: Pesticide / fertiliser application"),
              ("secF", "Section F: Harvesting"),
              ("total", "TOTAL cost of cultivation")]


    def stars(p):
        return "" if pd.isna(p) else "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.10 else ""


    def compare(s, incidence=True):
        t, c = s[T].dropna(), s[~T].dropna()
        tt = stats.ttest_ind(t, c, equal_var=False)
        mw = stats.mannwhitneyu(t, c, alternative="two-sided").pvalue
        sd_pool = np.sqrt((t.var(ddof=1) + c.var(ddof=1)) / 2)
        out = {"N_T": len(t), "Mean_T": t.mean(), "SD_T": t.std(ddof=1), "Median_T": t.median(),
               "N_C": len(c), "Mean_C": c.mean(), "SD_C": c.std(ddof=1), "Median_C": c.median(),
               "Diff": t.mean() - c.mean(), "p_t": tt.pvalue, "p_mw": mw,
               "NormDiff": (t.mean() - c.mean()) / sd_pool if sd_pool > 0 else np.nan}
        if incidence:
            out["Inc_T"], out["Inc_C"] = (t > 0).mean(), (c > 0).mean()
            tab = np.array([[(t > 0).sum(), (t == 0).sum()], [(c > 0).sum(), (c == 0).sum()]])
            out["p_inc"] = stats.fisher_exact(tab)[1] if tab.min() >= 0 and (tab[:, 1].sum() > 0) else np.nan
        out["Stars"] = stars(out["p_t"])
        return out


    rows = []
    for panel, keys in PANELS:
        rows.append({"Variable": panel, "_panel": True})
        for k in keys:
            s = d[f"cpa_{k}"].where(~flag[f"cpa_{k}"])
            note = ""
            rows.append({"Variable": LAB[k] + note, "var": f"cpa_{k}", "_panel": False, "n_excl": int(flag[f"cpa_{k}"].sum()), **compare(s)})
    rows.append({"Variable": "Section totals and total cost", "_panel": True})
    for k, lab in TOTALS:
        s = d[f"cpa_{k}"].where(~flag[f"cpa_{k}"])
        rows.append({"Variable": lab, "var": f"cpa_{k}", "_panel": False, "n_excl": int(flag[f"cpa_{k}"].sum()), **compare(s, incidence=False)})
    rows.append({"Variable": "Context", "_panel": True})
    rows.append({"Variable": "Rice area, Virippu 2025 (acres)", "var": "rice_acres_2025", "_panel": False, "n_excl": 0,
                 **compare(acres.where(acres > 0), incidence=False)})
    s_tot = d["cost_total_x"].where(~flag["cpa_total_x"] & (acres > 0))
    rows.append({"Variable": "Total cost (Rs., not per acre)", "var": "cost_total_x", "_panel": False,
                 "n_excl": int(flag["cpa_total_x"].sum()), **compare(s_tot, incidence=False)})
    res = pd.DataFrame(rows)

    # ---------------- cost shares (composition) ----------------
    ok_share = ~flag["cpa_total_x"] & ~flag[[f"cpa_sec{s}" for s in "ABDEF"] + ["cpa_secC_x"]].any(axis=1) & (d["cost_total_x"] > 0)
    share_rows = []
    for k, lab in [("secA", "A. Pre-nursery preparation"), ("secB", "B. Nursery preparation"), ("secC_x", "C. Transplanting"),
                   ("secD", "D. Weedicide / fertiliser application"), ("secE", "E. Pesticide / fertiliser application"), ("secF", "F. Harvesting")]:
        sh = (d[f"cost_{k}"] / d["cost_total_x"]).where(ok_share)
        r = compare(sh, incidence=False)
        share_rows.append({"Section": lab, **r})
    # labour / material / machinery split
    GROUPS = {
        "Hired labour": ["b_labour", "c_mainlab", "c_translab", "c_fert1lab", "c_gapfill", "d_weedlab", "d_fert2lab", "e_pestlab", "e_fert3lab", "f_labour"],
        "Machinery & transport": ["a_land", "c_mach", "f_mach", "f_tractor"],
        "Seed, fertiliser & plant protection": ["a_seed", "b_fert", "d_weed", "d_fert2", "e_pest", "e_fert3"],
        "Other items (mixed)": ["a_other", "b_other", "c_other", "d_other", "e_other", "f_other"],
    }
    for g, ks in GROUPS.items():
        sh = (d[[f"cost_{k}" for k in ks]].sum(axis=1) / d["cost_total_x"]).where(ok_share)
        share_rows.append({"Section": g, **compare(sh, incidence=False)})
    shares = pd.DataFrame(share_rows)

    # ---------------- robustness ----------------
    rob = []
    for k, lab in [(k, LAB[k]) for k in C.ORDER] + TOTALS:
        full = d[f"cpa_{k}"]
        newf = full.where(~old & ~flag[f"cpa_{k}"])
        a = compare(full, incidence=False); c3 = compare(newf, incidence=False)
        rob.append({"Variable": lab, "var": f"cpa_{k}",
                    "All obs: Mean T": a["Mean_T"], "All obs: Mean C": a["Mean_C"], "All obs: Median T": a["Median_T"], "All obs: Median C": a["Median_C"],
                    "All obs: p (t)": a["p_t"], "All obs: p (MW)": a["p_mw"],
                    "Newer form only: N T/C": f"{c3['N_T']}/{c3['N_C']}", "Newer form only: Diff": c3["Diff"], "Newer form only: p (t)": c3["p_t"]})
    rob = pd.DataFrame(rob)

    # ---------------- write ----------------
    nT, nC = int(T.sum()), int((~T).sum())
    hdr = ["Cost per acre (Rs./acre)", "N", "Mean", "SD", "Median", "% incurring", "N", "Mean", "SD", "Median", "% incurring",
           "Difference (T − C)", "", "p-value (t-test)", "p-value (Mann–Whitney)", "Normalized difference", "p-value (% incurring)", "Outliers excluded"]
    out_rows = []
    for _, r in res.iterrows():
        if r["_panel"]:
            out_rows.append([r["Variable"]] + [None] * 17)
        else:
            out_rows.append([r["Variable"], r["N_T"], r["Mean_T"], r["SD_T"], r["Median_T"], r.get("Inc_T"),
                             r["N_C"], r["Mean_C"], r["SD_C"], r["Median_C"], r.get("Inc_C"),
                             r["Diff"], r["Stars"], r["p_t"], r["p_mw"], r["NormDiff"], r.get("p_inc"), r["n_excl"]])
    sh_hdr = ["Share of total cost", "N T", "Mean share T", "N C", "Mean share C", "Difference (T − C)", "", "p-value (t-test)", "p-value (Mann–Whitney)"]
    sh_rows = [[r["Section"], r["N_T"], r["Mean_T"], r["N_C"], r["Mean_C"], r["Diff"], r["Stars"], r["p_t"], r["p_mw"]] for _, r in shares.iterrows()]

    act = res[(~res["_panel"]) & res["var"].astype(str).str.match(r"cpa_[a-f]_")]
    n10 = int((act["p_t"] < 0.10).sum())
    notes = [
        ["Groups", f"Treatment: payment = 1 (N = {nT}). Control: payment = 0 (N = {nC}), incl. 5 listed farmers with Payment = 'No'."],
        ["Timing", "Costs refer to Virippu 2025 (last season), before the RCT payments, so differences are baseline differences — not effects of the payment."],
        ["Measure", "Cost per acre = activity cost ÷ rice_acres_2025, from the cost variables in Survey_Round_1_merged_wide.xlsx. "
                    "Households that did not incur a cost are included as 0. Households with rice area 0 (3 control, 4 treatment) have no per-acre values."],
        ["Outliers", "Values flagged in the cost_outliers sheet (|modified z| > 3.5 on log scale) are excluded from that row only; the count is in the last column. "
                     "The robustness sheet repeats the tests on all observations (where medians and the Mann–Whitney test are the meaningful comparison, "
                     "because the entry errors make means meaningless) and on newer-form-version households only."],
        ["Tests", "Welch two-sample t-test on means; Mann–Whitney rank-sum test (robust to skew); normalized difference = diff ÷ √[(SD²_T + SD²_C)/2]; "
                  "Fisher's exact test for the share of households incurring the cost. *** p<0.01, ** p<0.05, * p<0.10."],
        ["Form versions", f"The older form version ({int(old.sum())} households: {int((old & T).sum())} treatment, {int((old & ~T).sum())} control) recorded 1st-dose "
                          "fertiliser as one item in the main sheet; the newer version allowed up to 2 items. The robustness sheet repeats the tests on newer-version households only."],
        ["Multiple comparisons", f"27 activities are tested; at the 10% level about 3 significant differences are expected by chance. Observed: {n10}."],
        ["Shares", "Share of each section in the household's total cost, for households with no flagged total or section."],
    ]

    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        pd.DataFrame(out_rows, columns=hdr).to_excel(xw, sheet_name="cost_comparison", index=False, startrow=3)
        pd.DataFrame(sh_rows, columns=sh_hdr).to_excel(xw, sheet_name="cost_shares", index=False, startrow=2)
        rob.drop(columns="var").to_excel(xw, sheet_name="robustness", index=False, startrow=2)
        pd.DataFrame(notes, columns=["Item", "Note"]).to_excel(xw, sheet_name="notes", index=False)

    wb = load_workbook(OUT)
    F = lambda **k: Font(**{"name": "Arial", "size": 10, **k})
    thin = Side(style="thin", color="000000")
    hfill = PatternFill("solid", start_color="DDE4EE"); pfill = PatternFill("solid", start_color="F2F2F2")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = F()

    ws = wb["cost_comparison"]
    ws["A1"] = "Table: Cost of rice cultivation per acre (Virippu 2025) by payment status"
    ws["A1"].font = F(bold=True, size=12)
    ws["A2"] = (f"Treatment = payment received (N = {nT}); Control = no payment (N = {nC}). Flagged outliers excluded. "
                "Welch t-test; *** p<0.01, ** p<0.05, * p<0.10. See 'notes'.")
    ws["A2"].font = F(italic=True)
    ws["B3"] = "Treatment (payment = 1)"; ws.merge_cells("B3:F3")
    ws["G3"] = "Control (payment = 0)"; ws.merge_cells("G3:K3")
    ws["L3"] = "Comparison"; ws.merge_cells("L3:R3")
    ws.merge_cells("L4:M4")
    for cell in list(ws[3]) + list(ws[4]):
        cell.font = F(bold=True); cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[4].height = 42
    for c in range(1, 19):
        ws.cell(3, c).border = Border(top=thin); ws.cell(4, c).border = Border(bottom=thin)
    for i, (_, r) in enumerate(res.iterrows()):
        rr = 5 + i
        if r["_panel"]:
            ws.cell(rr, 1).font = F(bold=True)
            for c in range(1, 19):
                ws.cell(rr, c).fill = pfill
            continue
        ws.cell(rr, 1).alignment = Alignment(indent=1)
        acre_row = r["var"] == "rice_acres_2025"
        for c in (3, 4, 5, 8, 9, 10, 12):
            ws.cell(rr, c).number_format = "0.00" if acre_row else "#,##0"
        for c in (6, 11):
            ws.cell(rr, c).number_format = "0%"
        for c in (14, 15, 16, 17):
            ws.cell(rr, c).number_format = "0.000"
        if r["p_t"] < 0.10:
            for c in (12, 13, 14):
                ws.cell(rr, c).font = F(bold=True)
        if pd.notna(r.get("p_mw")) and r["p_mw"] < 0.10:
            ws.cell(rr, 15).font = F(bold=True)
        if pd.notna(r.get("p_inc")) and r["p_inc"] < 0.10:
            ws.cell(rr, 17).font = F(bold=True)
        if r["var"] in ("cpa_total",):
            for c in range(1, 19):
                ws.cell(rr, c).font = F(bold=True)
    last = 5 + len(res) - 1
    for c in range(1, 19):
        ws.cell(last, c).border = Border(bottom=thin)
    ws.cell(last + 2, 1, f"Activities with a significant difference in means at 10%: {n10} of 27 (about 3 expected by chance). "
                         "Bold p-values are below 0.10.").font = F(italic=True)
    for j, w in enumerate([68, 5, 10, 10, 10, 9, 5, 10, 10, 10, 9, 12, 4, 10, 12, 11, 11, 9], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B5"
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.print_title_rows = "3:4"

    ws = wb["cost_shares"]
    ws["A1"] = "Composition of cultivation cost: mean share of total cost by section and by input type"
    ws["A1"].font = F(bold=True, size=12)
    for cell in ws[3]:
        cell.font = F(bold=True); cell.fill = hfill; cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[3].height = 32
    for row in ws.iter_rows(min_row=4):
        for c in (3, 5, 6):
            row[c - 1].number_format = "0.0%"
        for c in (8, 9):
            row[c - 1].number_format = "0.000"
    ws.insert_rows(10)
    ws.cell(10, 1, "By input type").font = F(bold=True)
    for j, w in enumerate([46, 6, 13, 6, 13, 13, 4, 11, 13], 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    ws = wb["robustness"]
    ws["A1"] = "Robustness: all observations incl. outliers (compare medians and Mann–Whitney; means are distorted by entry errors), and newer form version only (outliers excluded)"
    ws["A1"].font = F(bold=True, size=12)
    for cell in ws[3]:
        cell.font = F(bold=True); cell.fill = hfill; cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[3].height = 45
    for row in ws.iter_rows(min_row=4):
        for cell in row[1:]:
            if isinstance(cell.value, float):
                cell.number_format = "0.000" if "p (" in str(ws.cell(3, cell.column).value) else "#,##0"
    ws.column_dimensions["A"].width = 58
    for j in range(2, 11):
        ws.column_dimensions[get_column_letter(j)].width = 12

    ws = wb["notes"]
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = F(bold=True); cell.fill = hfill
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 125
    wb.save(OUT)

    pd.set_option("display.width", 260); pd.set_option("display.max_rows", 100)
    print(res[~res["_panel"]][["Variable", "N_T", "Mean_T", "Median_T", "N_C", "Mean_C", "Median_C", "Diff", "p_t", "p_mw", "Inc_T", "Inc_C", "p_inc", "n_excl"]].round(3).to_string(index=False))
    print(shares[["Section", "Mean_T", "Mean_C", "Diff", "p_t", "p_mw"]].round(3).to_string(index=False))
    print(rob[rob["var"].isin(["cpa_total", "cpa_c_fert1", "cpa_c_mach"])].round(3).to_string(index=False))
    print("n10 activities:", n10)


# =====================================================================================
# STEP 4 — item-level costs with dates, weekly aggregates  ->  items.csv, timeline_weekly.json
# =====================================================================================
def timeline_items(DATA, OUT):
    d = pd.read_excel(DATA, sheet_name="data")
    outl = pd.read_excel(DATA, sheet_name="cost_outliers", header=2)

    DATES = {
        "a_seed": ("a_seed_actdate", "a_seed_paydate"), "a_land": ("a_land_actdate", "a_land_paydate"),
        "b_labour": ("b_labour_actdate", "b_labour_paydate"), "b_fert": ("b_fert_actdate", "b_fert_paydate"),
        "c_mainlab": ("c_labour_actdate_land", "c_labour_paydate_land"), "c_mach": ("c_mach_actdate_land", "c_mach_paydate_land"),
        "c_translab": ("c_labour_actdate", "c_labour_paydate"), "c_fert1lab": ("c_labour_fert1_actdate", "c_labour_fert1_paydate"),
        "c_gapfill": ("c_labour_gapfilling_actdate", "c_labour_gapfilling_paydate"),
        "d_weed": ("d_weed_actdate", "d_weed_paydate"), "d_weedlab": ("d_weedlab_actdate", "d_weedlab_paydate"),
        "d_fert2lab": ("d_labour_fert2_actdate", "d_labour_fert2_paydate"),
        "e_pest": ("e_pest_actdate", "e_pest_paydate"), "e_pestlab": ("e_labour_pest_actdate", "e_labour_pest_paydate"),
        "e_fert3lab": ("e_labour_fert3_actdate", "e_labour_fert3_paydate"),
        "f_mach": ("f_mach_actdate", "f_mach_paydate"), "f_tractor": ("f_tractor_actdate", "f_tractor_paydate"),
        "f_labour": ("f_labour_actdate", "f_labour_paydate"),
    }

    # high outliers (entry errors) excluded per household × activity
    bad = set(zip(outl.loc[outl["direction"] == "high", "_index"], outl.loc[outl["direction"] == "high", "variable"].str[4:]))

    recs = []
    def add(key, label, sec, cost, act, pay):
        for i in cost.index:
            recs.append({"_index": d.at[i, "_index"], "farmers_id": d.at[i, "farmers_id"], "payment": d.at[i, "payment"],
                         "key": key, "label": label, "sec": sec, "cost": cost.at[i], "act": act.at[i], "pay": pay.at[i]})

    for key, label, sec, q, u, p in C.FIXED:
        c, _ = C.unit_price_cost(d[q], d[u], d[p])
        a, pdt = DATES[key]
        add(key, label, sec, c, pd.to_datetime(d[a], errors="coerce"), pd.to_datetime(d[pdt], errors="coerce"))
    for key, label, sec, pre, n in C.FERT:
        for q, u, p, nm in C.fert_items(d, pre, n):          # includes the older form's single 1st-dose item (c_fert1_qty ...)
            suffix = q[len(f"{pre}_qty"):]                     # "_1", "_2", ... or "" for the main-sheet item
            c, _ = C.unit_price_cost(d[q], d[u], d[p])
            add(key, label, sec, c, pd.to_datetime(d[f"{pre}_actdate{suffix}"], errors="coerce"), pd.to_datetime(d[f"{pre}_paydate{suffix}"], errors="coerce"))
    for key, label, sec, pre, n in C.OTHER:
        for k in range(1, n + 1):
            c, _ = C.other_cost(d[f"{pre}_qty_{k}"], d[f"{pre}_unit_{k}"], d[f"{pre}_amt_{k}"], d[f"{pre}_item_{k}"])
            add(key, label, sec, c, pd.to_datetime(d[f"{pre}_actdate_{k}"], errors="coerce"), pd.to_datetime(d[f"{pre}_paydate_{k}"], errors="coerce"))

    r = pd.DataFrame(recs)
    r = r[(r["cost"] > 0)]
    n_items = len(r)
    r["bad"] = [(i, k) in bad for i, k in zip(r["_index"], r["key"])]
    n_bad = int(r["bad"].sum()); bad_amt = float(r.loc[r["bad"], "cost"].sum())
    r = r[~r["bad"]]

    # dates: most were saved with year 2026 although they refer to Virippu 2025 → put every date on the 2025 calendar
    log = {}
    for col in ("act", "pay"):
        s = r[col]
        log[f"{col}_year_counts"] = {str(int(k)): int(v) for k, v in s.dt.year.value_counts().items()}
        norm = pd.to_datetime({"year": 2025, "month": s.dt.month, "day": s.dt.day}, errors="coerce").where(s.notna())
        placeholder = (s.dt.month == 1) & (s.dt.day == 1)
        out_window = norm < pd.Timestamp("2025-02-01")
        log[f"{col}_dropped_placeholder_or_january"] = int((placeholder | out_window).sum())
        r[col + "_n"] = norm.where(~(placeholder | out_window))
        r[col + "_wk"] = r[col + "_n"] - pd.to_timedelta(r[col + "_n"].dt.weekday, unit="D")   # Monday of the week

    r.drop(columns=["bad"]).to_csv(os.path.join(os.path.dirname(OUT), "items.csv"), index=False)   # item-level records for static plots
    ORDER = C.ORDER
    LAB = {k: l for k, l, *_ in C.FIXED} | {k: l for k, l, *_ in C.FERT} | {k: l for k, l, *_ in C.OTHER}
    SEC = {k: s for k, _, s, *_ in C.FIXED} | {k: s for k, _, s, *_ in C.FERT} | {k: s for k, _, s, *_ in C.OTHER}

    out = {"activities": [{"key": k, "label": LAB[k], "sec": SEC[k]} for k in ORDER],
           "sections": [{"sec": s, "label": l} for s, l in C.SECTIONS.items()],
           "cells": {}, "meta": {}}
    for col, name in (("act", "activity"), ("pay", "payment")):
        g = r.dropna(subset=[col + "_wk"]).groupby(["key", col + "_wk"])
        cells = g.agg(farmers=("_index", "nunique"), total=("cost", "sum"), items=("cost", "size")).reset_index()
        cells["week"] = cells[col + "_wk"].dt.strftime("%Y-%m-%d")
        out["cells"][name] = cells[["key", "week", "farmers", "total", "items"]].round({"total": 0}).to_dict("records")
        out["meta"][name] = {"items": int(r[col + "_wk"].notna().sum()), "amount": float(r.loc[r[col + "_wk"].notna(), "cost"].sum()),
                             "households": int(r.loc[r[col + "_wk"].notna(), "_index"].nunique()),
                             "undated_items": int(r[col + "_wk"].isna().sum()), "undated_amount": float(r.loc[r[col + "_wk"].isna(), "cost"].sum())}
        sw = r.dropna(subset=[col + "_wk"]).groupby(["sec", col + "_wk"]).agg(farmers=("_index", "nunique"), total=("cost", "sum")).reset_index()
        sw["week"] = sw[col + "_wk"].dt.strftime("%Y-%m-%d")
        out.setdefault("secweeks", {})[name] = sw[["sec", "week", "farmers", "total"]].round({"total": 0}).to_dict("records")
    out["meta"].update({"items_with_cost": n_items, "excluded_outlier_items": n_bad, "excluded_outlier_amount": bad_amt,
                        "households_total": int(d.shape[0]), **log,
                        "households_all_dates_2026": int((pd.concat([pd.to_datetime(d[c], errors='coerce').dt.year for c in d.columns if c.endswith(('_actdate', '_actdate_land')) or '_actdate_' in c], axis=1)
                                                          .apply(lambda row: (row.dropna() == 2026).all() and row.notna().any(), axis=1)).sum())})
    json.dump(out, open(OUT, "w"), default=str)
    print(json.dumps(out["meta"], indent=1, default=str))
    for name in ("activity", "payment"):
        c = pd.DataFrame(out["cells"][name])
        print(name, "cells", len(c), "weeks", c["week"].nunique(), c["week"].min(), c["week"].max(), "max farmers", c["farmers"].max(), "max total", c["total"].max())


# =====================================================================================
# STEP 5 — bubble plot of weekly cost per acre with cumulative line  ->  JPEG
# =====================================================================================
def bubble_plot(ITEMS, DATA, OUT, BASIS="act", MODE="overlay"):
    """BASIS: 'act' (activity date) or 'pay' (payment date). MODE: 'overlay' (line on the bubble axis) or 'panels'."""
    it = pd.read_csv(ITEMS, parse_dates=[f"{BASIS}_wk"])
    hh = pd.read_excel(DATA, sheet_name="data", usecols=["_index", "rice_acres_2025"])
    acres = hh.set_index("_index")["rice_acres_2025"]
    valid = acres[acres > 0]
    it = it[it["_index"].isin(valid.index)].dropna(subset=[f"{BASIS}_wk"]).copy()
    it["cpa"] = it["cost"] / it["_index"].map(acres)
    N_HH = len(valid)

    # bubbles: per farmer, per activity, per week → cost per acre; then average over the farmers who spent that week
    fw = it.groupby(["key", "sec", "label", "_index", f"{BASIS}_wk"], as_index=False)["cpa"].sum()
    bub = fw.groupby(["key", "sec", "label", f"{BASIS}_wk"]).agg(farmers=("_index", "nunique"), cpa=("cpa", "mean")).reset_index()
    bub["x"] = bub[f"{BASIS}_wk"] + pd.Timedelta(days=3.5)

    # cumulative line: average over all households with rice area of cumulative cost per acre
    wk = it.groupby(f"{BASIS}_wk")["cpa"].sum().sort_index() / N_HH
    weeks = pd.date_range("2025-01-27", "2025-12-29", freq="W-MON")
    cum = wk.reindex(weeks, fill_value=0).cumsum()
    cum = cum[(cum.index >= wk.index.min() - pd.Timedelta(days=7)) & (cum.index <= wk.index.max())]
    cx = cum.index + pd.Timedelta(days=7)          # value reached by the end of the week
    end_val = cum.iloc[-1]
    half_wk = cum.index[np.argmax(cum.values >= end_val / 2)]

    # ---------------- figure ----------------
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10.5})
    INK, INK2, MUTED, GRID, AXIS, SURF = "#141613", "#4d524a", "#868a80", "#e2e4dc", "#c3c6ba", "#fcfcfb"
    COL = {"A": "#2a78d6", "B": "#eb6834", "C": "#1baf7a", "D": "#eda100", "E": "#e87ba4", "F": "#008300"}
    SECLAB = {"A": "A. Pre-nursery preparation", "B": "B. Nursery preparation", "C": "C. Transplanting",
              "D": "D. Weedicide / 2nd fertiliser dose", "E": "E. Pesticide / 3rd fertiliser dose", "F": "F. Harvesting"}
    basis_word = "activity" if BASIS == "act" else "payment"
    inr = lambda v: f"₹{v:,.0f}"

    def style(ax, xlabel=True):
        ax.set_facecolor(SURF)
        ax.set_xlim(pd.Timestamp("2025-02-01"), pd.Timestamp("2025-12-31"))
        ax.xaxis.set_major_locator(mdates.MonthLocator())
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%b"))
        ax.grid(axis="y", color=GRID, lw=0.8, zorder=0); ax.grid(axis="x", color=GRID, lw=0.6, zorder=0)
        for sp in ("top", "right", "left"):
            ax.spines[sp].set_visible(False)
        ax.spines["bottom"].set_color(AXIS)
        ax.tick_params(colors=MUTED, length=0, labelsize=10)
        if xlabel:
            ax.set_xlabel(f"Week of {basis_word}, 2025", color=INK2)

    panels = MODE == "panels"
    if panels:
        fig, (ax, ax2) = plt.subplots(2, 1, figsize=(13, 9.4), dpi=200, sharex=True, gridspec_kw={"height_ratios": [3, 1.35], "hspace": 0.12})
        fig.subplots_adjust(left=0.085, right=0.78, top=0.865, bottom=0.135)
    else:
        fig, ax = plt.subplots(figsize=(13, 7.6), dpi=200)
        fig.subplots_adjust(left=0.085, right=0.78, top=0.82, bottom=0.185)
    fig.patch.set_facecolor("white")

    maxf = bub["farmers"].max()
    size = lambda f: (np.sqrt(f / maxf) * 26) ** 2          # marker area proportional to farmers
    b = bub.sort_values("farmers", ascending=False)
    ax.scatter(b["x"], b["cpa"], s=size(b["farmers"]), c=b["sec"].map(COL), alpha=0.78, edgecolors="white", linewidths=0.7, zorder=3)
    ax.set_yscale("log")
    yt = [10, 50, 100, 500, 1000, 5000, 10000, 50000]
    lo = max(5, bub["cpa"].min() * 0.6)
    hi = (bub["cpa"].max() if panels else max(end_val, bub["cpa"].max())) * 1.9
    ax.set_ylim(lo, hi)
    ax.yaxis.set_major_locator(FixedLocator([t for t in yt if lo <= t <= hi]))
    ax.yaxis.set_minor_locator(NullLocator())
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: inr(v)))
    style(ax, xlabel=not panels)
    ax.set_ylabel("Cost per acre (₹, log scale)", color=INK2)

    hx = half_wk + pd.Timedelta(days=7); hy = cum.loc[half_wk]
    half_txt = f"Half the season's cost\nspent by {hx - pd.Timedelta(days=1):%-d %b}"
    end_txt = f"{inr(end_val)} per acre\nby season end"
    if panels:
        ax2.fill_between(cx, 0, cum.values, color=INK, alpha=0.07, lw=0, zorder=2)
        ax2.plot(cx, cum.values, color=INK, lw=2.2, zorder=4, solid_capstyle="round")
        ax2.set_ylim(0, end_val * 1.18)
        ax2.yaxis.set_major_formatter(FuncFormatter(lambda v, _: inr(v)))
        ax2.set_yticks([0, 10000, 20000, 30000, 40000])
        style(ax2)
        ax2.set_ylabel("Cumulative cost\nper acre (₹)", color=INK2)
        tgt = ax2
    else:
        m = cum.values >= lo
        ax.plot(cx[m], cum.values[m], color=INK, lw=2.2, zorder=4, solid_capstyle="round")
        tgt = ax
    tgt.scatter([cx[-1]], [end_val], s=28, color=INK, zorder=5)
    tgt.annotate(end_txt, (cx[-1], end_val), xytext=(-8, -22) if panels else (-10, -10), textcoords="offset points",
                 va="top", ha="right", fontsize=10, color=INK, fontweight="bold")
    tgt.scatter([hx], [hy], s=26, color=INK, zorder=5)
    tgt.annotate(half_txt, (hx, hy), xytext=(12, -4), textcoords="offset points", ha="left", va="top", fontsize=9.5, color=INK2)

    title = "Rice cultivation cost per acre by week and activity, Virippu 2025"
    sub = ("Bubbles: average cost per acre among farmers who spent on that activity in the week; bubble size is the number of farmers. "
           + ("Lower panel" if panels else "Black line") + f": cumulative cost per acre of the average farmer ({N_HH} households).")
    fig.text(0.085, 0.955 if panels else 0.945, title, fontsize=16, fontweight="bold", color=INK)
    fig.text(0.085, 0.94 if panels else 0.928, "\n".join(textwrap.wrap(sub, 150)), fontsize=10.5, color=INK2, va="top")

    h_sec = [Line2D([0], [0], marker="o", ls="", markersize=9, markerfacecolor=COL[k], markeredgecolor="white", label=SECLAB[k]) for k in COL]
    top = 0.865 if panels else 0.82
    leg1 = fig.legend(handles=h_sec, title="Stage of cultivation", loc="upper left", bbox_to_anchor=(0.79, top), frameon=False,
                      fontsize=9.5, title_fontsize=10, alignment="left")
    h_sz = [plt.scatter([], [], s=size(v), facecolor="none", edgecolor=MUTED, lw=1, label=f"{v} farmers") for v in (10, 50, 100)]
    leg2 = fig.legend(handles=h_sz, title="Farmers in the week", loc="upper left", bbox_to_anchor=(0.79, top - (0.25 if panels else 0.30)), frameon=False,
                      fontsize=9.5, title_fontsize=10, labelspacing=1.5, borderpad=0.8, alignment="left")
    if not panels:
        fig.legend(handles=[Line2D([0], [0], color=INK, lw=2.2, label="Cumulative cost per acre,\naverage farmer")],
                   loc="upper left", bbox_to_anchor=(0.79, top - 0.56), frameon=False, fontsize=9.5)
    for lg in (leg1, leg2):
        lg.get_title().set_color(INK)

    note = (f"Source: Survey Round 1, Section 3 (Cash transfer RCT); {N_HH} households with rice area > 0. Cost = quantity × price per unit; "
            f"'other items' = amount paid. Weeks start Monday. Dates are {basis_word} dates placed on the 2025 calendar (246 households recorded 2026 by mistake). "
            "Items flagged as entry errors and items without a date are excluded." + ("" if panels else " The y-axis is logarithmic, which compresses the later rise of the cumulative line."))
    fig.text(0.085, 0.025, "\n".join(textwrap.wrap(note, 180)), fontsize=8.3, color=MUTED, va="bottom")
    fig.savefig(OUT, dpi=200, pil_kwargs={"quality": 95})
    print(f"{MODE}: end {end_val:,.0f}; half by week of {half_wk:%Y-%m-%d}; bubbles {len(bub)}; households {N_HH}")
    plt.close(fig)


# =====================================================================================
# RUN ALL STEPS
# =====================================================================================
def main(data_dir=DATA_DIR, out_dir=OUT_DIR):
    data_dir, out_dir = Path(data_dir), Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    merged = out_dir / "Survey_Round_1_merged_wide.xlsx"
    print("Step 1: merging sheets and computing costs");      merge_survey(data_dir, merged)
    print("Step 2: balance table");                           balance_table(merged, out_dir / "Balance_table_payment_round1.xlsx")
    print("Step 3: cost comparison");                         cost_comparison(merged, out_dir / "Cost_comparison_payment_round1.xlsx")
    print("Step 4: weekly timeline data");                    timeline_items(merged, out_dir / "timeline_weekly.json")
    print("Step 5: bubble plots")
    for mode in ("overlay", "panels"):
        bubble_plot(out_dir / "items.csv", merged, out_dir / f"Cost_per_acre_bubble_{mode}.jpg", "act", mode)
    print(f"Done. Outputs in {out_dir.resolve()}")


if __name__ == "__main__":
    if len(sys.argv) == 3:
        main(sys.argv[1], sys.argv[2])
    else:
        main()
